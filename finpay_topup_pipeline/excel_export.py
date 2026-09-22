from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .config import (
    TABLE_MANUAL_ADJUSTMENT,
    TABLE_REFRESH,
    TABLE_REFRESH_CLUSTER,
    TABLE_TXN_STAGING,
)
from .saldo import _checkpoint_source_sql


LOCAL_TZ = ZoneInfo("Asia/Makassar")
MAX_EXPORT_DAYS = 366
MAX_EXPORT_ROWS = 200_000
MAX_TELEGRAM_FILE_BYTES = 50 * 1024 * 1024
RUPIAH_FORMAT = '_-"Rp"* #,##0_-;\\-"Rp"* #,##0_-;_-"Rp"* "-"??_-;_-@_-'
HEADERS = (
    "Transaction Date",
    "Sender",
    "Receiver",
    "Transaction Type",
    "SETOR",
    "TOPUP",
    "SALDO",
    "Currency",
    "Remarks",
    "Outlet",
    "Source",
)
STAGING_HEADERS = (
    "Refresh ID",
    "Staging ID",
    "Transaction Date",
    "Sender",
    "Receiver",
    "Transaction Type",
    "SETOR",
    "TOPUP",
    "SALDO",
    "Currency",
    "Remarks",
    "Source File",
    "Row Hash",
    "Source Status",
)
DAILY_HEADERS = (
    "Cluster",
    "Date",
    "Rows",
    "SETOR",
    "TOPUP",
    "Net",
    "First Transaction",
    "Last Transaction",
    "Coverage",
)
SUMMARY_HEADERS = (
    "Cluster",
    "Verification",
    "Checkpoint Boundary",
    "Download Period",
    "Opening SALDO",
    "Calculated SALDO",
    "CMS BUCKET",
    "Difference",
    "Staging Rows",
    "Days Without Rows",
)
CLUSTER_SHEETS = {
    "411311": "Palangkaraya",
    "421306": "Morotai",
    "421307": "Tidore",
    "421315": "Banggai",
    "421318": "Morowali",
    "421320": "Ternate",
}


def _as_date(value):
    return value if isinstance(value, date) and not isinstance(value, datetime) else date.fromisoformat(str(value)[:10])


def _local_date(value):
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    return value.astimezone(LOCAL_TZ).date() if value.tzinfo is not None else value.date()


def _excel_datetime(value):
    return value.astimezone(LOCAL_TZ).replace(tzinfo=None) if value.tzinfo is not None else value


def _cluster_rows(conn, cluster_id, start_date, end_date):
    with conn.cursor() as cur:
        cur.execute(
            f"SELECT opening_balance, as_of_date FROM ({_checkpoint_source_sql()}) checkpoints "
            "WHERE cluster_id=%s AND as_of_date<=%s "
            "ORDER BY as_of_date DESC, priority DESC, override_id DESC LIMIT 1",
            (cluster_id, start_date),
        )
        checkpoint = cur.fetchone()
        if checkpoint is None:
            raise ValueError(f"checkpoint before {start_date} is missing for cluster {cluster_id}")
        opening_balance, opening_date = checkpoint
        cur.execute(
            "SELECT u.txn_id,u.transaction_date,u.sender,u.receiver,u.transaction_type,u.amount,"
            "u.currency,u.remarks,u.source_file,u.row_hash,u.source_kind,"
            "COALESCE(o.label,c.outlet_code) AS outlet_label "
            "FROM ("
            "SELECT txn_id,cluster_id,transaction_date,sender,receiver,transaction_type,amount,currency,"
            "remarks,source_file,row_hash,'DigiPOS' AS source_kind FROM finpay_topup_txn WHERE cluster_id=%(cluster_id)s "
            "UNION ALL "
            f"SELECT (-adjustment_id)::bigint,cluster_id,effective_at,sender,receiver,transaction_type,amount,currency,"
            f"remarks,'manual_adjustment',adjustment_hash,'Manual Adjustment' FROM {TABLE_MANUAL_ADJUSTMENT} "
            "WHERE cluster_id=%(cluster_id)s AND status='APPROVED'"
            ") u LEFT JOIN finpay_topup_classification c ON c.txn_id=u.txn_id AND u.source_kind='DigiPOS' "
            "LEFT JOIN finpay_outlet o ON o.code=c.outlet_code "
            "WHERE (u.transaction_date AT TIME ZONE 'Asia/Makassar')::date >= %(opening_date)s "
            "AND (u.transaction_date AT TIME ZONE 'Asia/Makassar')::date <= %(end_date)s "
            "ORDER BY u.transaction_date ASC",
            {"cluster_id": cluster_id, "opening_date": opening_date, "end_date": end_date},
        )
        source_rows = cur.fetchall()

    saldo = Decimal(str(opening_balance))
    export_rows = []
    for row in source_rows:
        amount = Decimal(str(row[5]))
        if row[4] == "Kredit":
            saldo += amount
        elif row[4] == "Debit":
            saldo -= amount
        if _local_date(row[1]) < start_date:
            continue
        export_rows.append({
            "transaction_date": row[1],
            "sender": row[2],
            "receiver": row[3],
            "transaction_type": row[4],
            "debit": amount if row[4] == "Debit" else None,
            "kredit": amount if row[4] == "Kredit" else None,
            "saldo": saldo,
            "currency": row[6] or "IDR",
            "remarks": row[7],
            "outlet": row[11],
            "source": row[10],
        })
    return export_rows


def build_finance_workbook(conn, start_date, end_date, cluster_ids):
    """Build a finance-style, read-only Top-Up workbook in memory."""
    start_date = _as_date(start_date)
    end_date = _as_date(end_date)
    if start_date > end_date:
        raise ValueError("start date must be on or before end date")
    if (end_date - start_date).days >= MAX_EXPORT_DAYS:
        raise ValueError(f"date range may not exceed {MAX_EXPORT_DAYS} days")
    cluster_ids = list(dict.fromkeys(str(cluster_id) for cluster_id in cluster_ids))
    if not cluster_ids or any(cluster_id not in CLUSTER_SHEETS for cluster_id in cluster_ids):
        raise ValueError("export cluster is not configured")

    workbook = Workbook()
    workbook.remove(workbook.active)
    counts = {}
    total_rows = 0
    for cluster_id in cluster_ids:
        rows = _cluster_rows(conn, cluster_id, start_date, end_date)
        total_rows += len(rows)
        if total_rows > MAX_EXPORT_ROWS:
            raise ValueError(f"export exceeds {MAX_EXPORT_ROWS:,} rows")
        counts[cluster_id] = len(rows)
        sheet = workbook.create_sheet(CLUSTER_SHEETS[cluster_id])
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append((
                _excel_datetime(row["transaction_date"]),
                row["sender"],
                row["receiver"],
                row["transaction_type"],
                row["kredit"],
                row["debit"],
                row["saldo"],
                row["currency"],
                row["remarks"],
                row["outlet"],
                row["source"],
            ))
        for cell in sheet["A"][1:]:
            cell.number_format = "dd-mm-yyyy hh:mm:ss"
        for column in ("E", "F", "G"):
            for cell in sheet[column][1:]:
                cell.number_format = RUPIAH_FORMAT
        widths = (20, 18, 18, 18, 16, 16, 18, 12, 32, 20, 20)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{max(sheet.max_row, 1)}"

    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    if len(content) > MAX_TELEGRAM_FILE_BYTES:
        raise ValueError("generated workbook exceeds Telegram's 50 MB document limit")
    scope = cluster_ids[0] if len(cluster_ids) == 1 else "all"
    return {
        "content": content,
        "filename": f"finpay-topup-{scope}-{start_date.isoformat()}-to-{end_date.isoformat()}.xlsx",
        "row_count": total_rows,
        "cluster_counts": counts,
        "start": start_date.isoformat(),
        "end": end_date.isoformat(),
    }


def build_staging_workbook(conn, refresh_id, cluster_ids=None):
    """Build a read-only analysis workbook for mismatched pending staging."""
    refresh_id = int(refresh_id)
    with conn.cursor() as cur:
        cur.execute(
            f"SELECT r.status, r.requested_start, r.requested_end, c.cluster_id, c.status, "
            f"c.checkpoint_as_of_date, c.computed_saldo, c.bucket_topup_value, "
            f"c.verification_diff, c.bucket_snapshot_at, c.calculation_cutoff_at "
            f"FROM {TABLE_REFRESH} r "
            f"JOIN {TABLE_REFRESH_CLUSTER} c ON c.refresh_id = r.refresh_id "
            "WHERE r.refresh_id = %s ORDER BY c.cluster_id",
            (refresh_id,),
        )
        refresh_rows = cur.fetchall()
    if not refresh_rows:
        raise ValueError("refresh was not found")
    if refresh_rows[0][0] != "PENDING_REVIEW":
        raise ValueError("staging export requires a pending review refresh")

    start_date = _as_date(refresh_rows[0][1])
    end_date = _as_date(refresh_rows[0][2])
    contexts = {
        str(row[3]): {
            "status": row[4],
            "checkpoint": _as_date(row[5]) if row[5] else None,
            "computed": row[6],
            "bucket": row[7],
            "difference": row[8],
            "snapshot_at": row[9],
            "cutoff_at": row[10],
        }
        for row in refresh_rows
    }
    mismatch_ids = [cluster_id for cluster_id, value in contexts.items() if value["status"] == "MISMATCH"]
    if cluster_ids is None:
        cluster_ids = mismatch_ids
    else:
        cluster_ids = list(dict.fromkeys(str(cluster_id) for cluster_id in cluster_ids))
        if any(cluster_id not in mismatch_ids for cluster_id in cluster_ids):
            raise ValueError("staging export is available only for mismatched clusters")
    if not cluster_ids:
        raise ValueError("no mismatched cluster requires a staging export")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"FinPay staging analysis refresh {refresh_id}"
    workbook.properties.subject = "Pending Top-Up staging review"
    workbook.properties.creator = "FinPay Top-Up"
    summary_rows = []
    daily_rows = []
    total_rows = 0
    gap_days = {}

    control = workbook.create_sheet("Ringkasan")
    control.append(("Refresh ID", refresh_id))
    control.append(("Status", "PENDING_REVIEW"))
    control.append(("Periode permintaan", f"{start_date.isoformat()} - {end_date.isoformat()}"))
    control.append(("Catatan", "Data di bawah hanya staging dan belum masuk ledger utama."))
    control.append(())
    control.append(SUMMARY_HEADERS)

    daily = workbook.create_sheet("Harian")
    daily.append(DAILY_HEADERS)

    for cluster_id in cluster_ids:
        context = contexts[cluster_id]
        checkpoint = context["checkpoint"]
        if checkpoint is None:
            raise ValueError(f"checkpoint is missing for cluster {cluster_id}")
        if (end_date - checkpoint).days >= MAX_EXPORT_DAYS:
            raise ValueError(f"staging period exceeds {MAX_EXPORT_DAYS} days")
        with conn.cursor() as cur:
            cur.execute(
                f"SELECT opening_balance FROM ({_checkpoint_source_sql()}) checkpoints "
                "WHERE cluster_id=%s AND as_of_date<=%s "
                "ORDER BY as_of_date DESC, priority DESC, override_id DESC LIMIT 1",
                (cluster_id, checkpoint),
            )
            opening = cur.fetchone()
            if opening is None:
                raise ValueError(f"checkpoint balance is missing for cluster {cluster_id}")
            cur.execute(
                f"SELECT txn_id, transaction_date, sender, receiver, transaction_type, amount, "
                f"currency, remarks, source_file, row_hash, loaded_at "
                f"FROM {TABLE_TXN_STAGING} WHERE refresh_id=%s AND cluster_id=%s "
                "ORDER BY transaction_date ASC, txn_id ASC",
                (refresh_id, cluster_id),
            )
            source_rows = cur.fetchall()

        saldo = Decimal(str(opening[0]))
        export_rows = []
        daily_by_date = {}
        for row in source_rows:
            transaction_day = _local_date(row[1])
            amount = Decimal(str(row[5]))
            if row[4] == "Kredit":
                saldo += amount
            elif row[4] == "Debit":
                saldo -= amount
            if transaction_day < checkpoint or transaction_day > end_date:
                continue
            item = daily_by_date.setdefault(
                transaction_day,
                {"rows": 0, "kredit": Decimal("0"), "debit": Decimal("0"), "first": row[1], "last": row[1]},
            )
            item["rows"] += 1
            item["kredit"] += amount if row[4] == "Kredit" else Decimal("0")
            item["debit"] += amount if row[4] == "Debit" else Decimal("0")
            item["first"] = min(item["first"], row[1])
            item["last"] = max(item["last"], row[1])
            export_rows.append((
                refresh_id,
                row[0],
                _excel_datetime(row[1]),
                row[2],
                row[3],
                row[4],
                amount if row[4] == "Kredit" else None,
                amount if row[4] == "Debit" else None,
                saldo,
                row[6] or "IDR",
                row[7],
                row[8],
                row[9],
                "STAGING - belum masuk ledger",
            ))

        total_rows += len(export_rows)
        if total_rows > MAX_EXPORT_ROWS:
            raise ValueError(f"export exceeds {MAX_EXPORT_ROWS:,} rows")
        missing_days = []
        day = checkpoint
        while day <= end_date:
            item = daily_by_date.get(day)
            if item is None:
                missing_days.append(day.isoformat())
                daily.append((CLUSTER_SHEETS[cluster_id], day, 0, None, None, None, None, None, "Tidak ada baris dikembalikan"))
            else:
                daily.append((
                    CLUSTER_SHEETS[cluster_id],
                    day,
                    item["rows"],
                    item["kredit"],
                    item["debit"],
                    item["kredit"] - item["debit"],
                    _excel_datetime(item["first"]),
                    _excel_datetime(item["last"]),
                    "Ada data",
                ))
            day += timedelta(days=1)
        gap_days[cluster_id] = missing_days
        summary_rows.append((
            CLUSTER_SHEETS[cluster_id],
            context["status"],
            checkpoint,
            f"{checkpoint.isoformat()} - {end_date.isoformat()}",
            Decimal(str(opening[0])),
            context["computed"],
            context["bucket"],
            context["difference"],
            len(export_rows),
            ", ".join(missing_days) or "Tidak ada",
        ))
        sheet = workbook.create_sheet(CLUSTER_SHEETS[cluster_id])
        sheet.append(STAGING_HEADERS)
        for row in export_rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:N{max(sheet.max_row, 1)}"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for cell in sheet["C"][1:]:
            cell.number_format = "dd-mm-yyyy hh:mm:ss"
        for column in ("G", "H", "I"):
            for cell in sheet[column][1:]:
                cell.number_format = RUPIAH_FORMAT
        widths = (12, 12, 20, 18, 18, 18, 16, 16, 18, 12, 32, 36, 68, 26)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + index)].width = width

    for row in summary_rows:
        control.append(row)
    for cell in control[6]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    control.freeze_panes = "A7"
    control.auto_filter.ref = f"A6:J{max(control.max_row, 6)}"
    for column, width in zip("ABCDEFGHIJ", (18, 16, 20, 28, 18, 18, 18, 18, 16, 38)):
        control.column_dimensions[column].width = width

    for cell in daily[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    daily.freeze_panes = "A2"
    daily.auto_filter.ref = f"A1:I{max(daily.max_row, 1)}"
    for cell in daily["B"][1:]:
        cell.number_format = "dd-mm-yyyy"
    for column in ("D", "E", "F"):
        for cell in daily[column][1:]:
            cell.number_format = RUPIAH_FORMAT
    for column, width in zip("ABCDEFGHI", (18, 14, 10, 16, 16, 16, 20, 20, 28)):
        daily.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    if len(content) > MAX_TELEGRAM_FILE_BYTES:
        raise ValueError("generated workbook exceeds Telegram's 50 MB document limit")
    return {
        "content": content,
        "filename": f"finpay-topup-staging-refresh-{refresh_id}-mismatch.xlsx",
        "refresh_id": refresh_id,
        "cluster_ids": cluster_ids,
        "row_count": total_rows,
        "cluster_counts": {row[0]: row[8] for row in summary_rows},
        "gap_days": gap_days,
        "start": start_date.isoformat(),
        "end": end_date.isoformat(),
    }
