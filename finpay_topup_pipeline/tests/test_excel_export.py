import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from tempfile import NamedTemporaryFile
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from finpay_topup_pipeline.excel_export import (
    HEADERS,
    STAGING_HEADERS,
    build_finance_workbook,
    build_staging_workbook,
)


class TestFinanceExcelExport(unittest.TestCase):
    def test_builds_finance_style_rows_with_running_saldo(self):
        conn = MagicMock()
        cur = conn.cursor.return_value.__enter__.return_value
        cur.fetchone.return_value = (Decimal("1000"), date(2026, 8, 20))
        cur.fetchall.return_value = [
            (
                1,
                datetime(2026, 8, 20, 1, tzinfo=timezone.utc),
                "sender",
                "receiver",
                "Kredit",
                Decimal("100"),
                "IDR",
                "setor",
                "source.csv",
                "hash-1",
                "DigiPOS",
                "BUNGKU",
            ),
            (
                2,
                datetime(2026, 8, 20, 2, tzinfo=timezone.utc),
                "sender",
                "receiver",
                "Debit",
                Decimal("50"),
                "IDR",
                "topup",
                "source.csv",
                "hash-2",
                "DigiPOS",
                None,
            ),
        ]

        result = build_finance_workbook(
            conn,
            "2026-08-20",
            "2026-08-20",
            ["421318"],
        )

        workbook = load_workbook(BytesIO(result["content"]), data_only=True)
        sheet = workbook["Morowali"]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
        self.assertEqual(sheet["E2"].value, 100)
        self.assertIsNone(sheet["F2"].value)
        self.assertEqual(sheet["G2"].value, 1100)
        self.assertEqual(sheet["F3"].value, 50)
        self.assertEqual(sheet["G3"].value, 1050)
        self.assertEqual(sheet["J2"].value, "BUNGKU")
        self.assertEqual(sheet["K2"].value, "DigiPOS")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:K3")
        self.assertEqual(result["row_count"], 2)
        query = cur.execute.call_args_list[1].args[0]
        self.assertIn("ORDER BY u.transaction_date ASC", query)
        self.assertNotIn("ORDER BY u.transaction_date ASC,", query)

    def test_rejects_unbounded_or_unknown_exports(self):
        with self.assertRaises(ValueError):
            build_finance_workbook(MagicMock(), "2025-01-01", "2026-01-02", ["421318"])
        with self.assertRaises(ValueError):
            build_finance_workbook(MagicMock(), "2026-01-01", "2026-01-02", ["999999"])

    def test_staging_workbook_shows_daily_gap_and_only_mismatch_cluster(self):
        conn = MagicMock()
        cur = conn.cursor.return_value.__enter__.return_value
        cur.fetchall.side_effect = [
            [
                (
                    "PENDING_REVIEW",
                    date(2026, 8, 11),
                    date(2026, 8, 13),
                    "421318",
                    "MISMATCH",
                    date(2026, 8, 11),
                    Decimal("1080"),
                    Decimal("1100"),
                    Decimal("-20"),
                    datetime(2026, 8, 14, tzinfo=timezone.utc),
                    datetime(2026, 8, 14, tzinfo=timezone.utc),
                )
            ],
            [
                (
                    1,
                    datetime(2026, 8, 11, 1, tzinfo=timezone.utc),
                    "sender",
                    "receiver",
                    "Kredit",
                    Decimal("100"),
                    "IDR",
                    "setor",
                    "source-11.csv",
                    "hash-11",
                    datetime(2026, 8, 14, tzinfo=timezone.utc),
                ),
                (
                    2,
                    datetime(2026, 8, 13, 1, tzinfo=timezone.utc),
                    "sender",
                    "receiver",
                    "Debit",
                    Decimal("20"),
                    "IDR",
                    "topup",
                    "source-13.csv",
                    "hash-13",
                    datetime(2026, 8, 14, tzinfo=timezone.utc),
                ),
            ],
        ]
        cur.fetchone.return_value = (Decimal("1000"),)

        result = build_staging_workbook(conn, 7)

        workbook = load_workbook(BytesIO(result["content"]), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Ringkasan", "Harian", "Morowali"])
        self.assertEqual(tuple(cell.value for cell in workbook["Morowali"][1]), STAGING_HEADERS)
        self.assertEqual(workbook["Morowali"]["N2"].value, "STAGING - belum masuk ledger")
        self.assertEqual(workbook["Morowali"].auto_filter.ref, "A1:N3")
        self.assertEqual(workbook["Ringkasan"].auto_filter.ref, "A6:J7")
        self.assertEqual(workbook["Harian"]["I3"].value, "Tidak ada baris dikembalikan")
        self.assertEqual(result["gap_days"], {"421318": ["2026-08-12"]})

    @patch("finpay_topup_pipeline.notify.requests.post")
    def test_send_telegram_document_uses_multipart_upload(self, post):
        from finpay_topup_pipeline.notify import send_telegram_document

        post.return_value.raise_for_status = MagicMock()
        with NamedTemporaryFile(suffix=".xlsx") as workbook:
            workbook.write(b"xlsx")
            workbook.flush()
            self.assertTrue(
                send_telegram_document(
                    workbook.name,
                    token="token",
                    chat_id="123",
                    filename="review.xlsx",
                    api_base="https://telegram.test",
                )
            )

        self.assertEqual(post.call_args.args[0], "https://telegram.test/bottoken/sendDocument")
        self.assertEqual(post.call_args.kwargs["data"], {"chat_id": "123"})
        self.assertEqual(post.call_args.kwargs["files"]["document"][0], "review.xlsx")
