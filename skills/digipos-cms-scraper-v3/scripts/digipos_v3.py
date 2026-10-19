#!/usr/bin/env python3
"""Standalone DigiPOS per-cell ingestion primitives for the v3 skill."""
from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import urllib.parse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

BASE_URL = os.environ.get("DIGIPOS_BASE_URL", "https://digipos-cms.finpay.id")
MONITORING_URL = f"{BASE_URL}/deposit/monitoring-riwayat"
DATA_ENDPOINT = re.compile(r"monitoring-riwayat-data")
DATA_ENDPOINT_PATH = "/deposit/monitoring-riwayat-data"
JQUERY_MIRROR = "https://cdn.jsdelivr.net/npm/jquery@3.6.4/dist/jquery.min.js"
MAX_PAGE_LENGTH = 20_000

# These are login identities, not cluster/cell IDs. Cell options are discovered
# from DigiPOS for each login rather than maintained in this account roster.
ACCOUNT_ROSTER = {
    "422040_B": "422040",
    "422140_B": "422140",
    "412110_B": "412110",
    "422010_B": "422010",
    "422090_B": "422090",
    "422100_B": "422100",
    "422180_B": "422180",
    "422220_B": "422220",
}

TITLE = "CMS DigiPOS - Riwayat Saldo"
HEADERS = [
    "No",
    "Transaction Date",
    "Transaction ID",
    "Saldo Awal",
    "Kredit",
    "Debet",
    "Saldo Akhir",
    "Transaction Type",
    "Transaction",
    "Nomor RS",
    "Remarks",
]
MONEY_COLUMNS = {4, 5, 6, 7}
WIDTHS = {
    "A": 6.0,
    "B": 25.65,
    "C": 37.8,
    "D": 14.85,
    "E": 13.5,
    "F": 14.85,
    "G": 14.85,
    "H": 21.6,
    "I": 18.9,
    "J": 27.0,
    "K": 54.0,
}

SELECT_CELL = """(value) => {
    const select = document.querySelector('#parentPhoneNo_pick');
    if (!select) return {ok: false, reason: 'cell selector missing'};
    const option = Array.from(select.options).find(item => item.value === value);
    if (!option) return {ok: false, reason: 'cell option missing'};
    select.value = value;
    jQuery(select).trigger('change');
    return {
        ok: true,
        label: (option.textContent || '').trim(),
        parentcode: (document.querySelector('#parentPhoneNo') || {}).value || null
    };
}"""

SET_DATES = """([displayStart, displayEnd, apiStart, apiEnd]) => {
    const setValue = (selector, value) => {
        const input = document.querySelector(selector);
        if (!input) throw new Error(`missing date field ${selector}`);
        const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value').set;
        setter.call(input, value);
        input.dispatchEvent(new Event('input', {bubbles: true}));
        input.dispatchEvent(new Event('change', {bubbles: true}));
    };
    setValue('input[name="date1"]', displayStart);
    setValue('input[name="date2"]', displayEnd);
    setValue('#dateStart', apiStart);
    setValue('#dateEnd', apiEnd);
}"""

READ_CELL_OPTIONS = """() => {
    const select = document.querySelector('#parentPhoneNo_pick');
    if (!select) return [];
    return Array.from(select.options).map(option => ({
        value: option.value,
        label: (option.textContent || '').trim()
    }));
}"""

READ_TABLE_STATE = """() => {
    const table = $('#dataInput').DataTable();
    const info = table.page.info();
    const text = ($('#dataInput_info').text() || '').trim();
    return {
        info_text: text,
        total: Number.isFinite(info.recordsDisplay) ? info.recordsDisplay : null,
        loading: $('#loadingDialog').is(':visible'),
        modal: $('#modal-gagal').is(':visible'),
        modal_text: ($('#modal-gagal').text() || '').trim().slice(0, 200),
        processing: Boolean(table.processing && table.processing())
    };
}"""


@dataclass(frozen=True)
class CellTarget:
    account: str
    cluster_id: str
    cell_code: str
    parentcode: str
    label: str


def cluster_id_for(account: str) -> str:
    try:
        return ACCOUNT_ROSTER[account]
    except KeyError as exc:
        raise ValueError(f"account is not in the v3 _B roster: {account}") from exc


def discover_cells(account: str, options: list[dict]) -> tuple[list[CellTarget], list[dict]]:
    """Return explicit district cells and separately record aggregate options."""
    cluster_id = cluster_id_for(account)
    targets: list[CellTarget] = []
    aggregates: list[dict] = []
    for option in options:
        parentcode = str(option.get("value") or "")
        label = str(option.get("label") or "").strip()
        if label.casefold().startswith("semua cell") or parentcode == cluster_id:
            aggregates.append({"parentcode": parentcode, "label": label})
            continue
        if not parentcode.startswith(cluster_id) or len(parentcode) <= len(cluster_id):
            raise ValueError(f"unexpected cell value for {account}: {parentcode!r}")
        targets.append(
            CellTarget(
                account=account,
                cluster_id=cluster_id,
                cell_code=parentcode[len(cluster_id) :],
                parentcode=parentcode,
                label=label,
            )
        )
    if not 1 <= len(targets) <= 3:
        raise ValueError(f"expected 1-3 selectable cells for {account}, found {len(targets)}")
    return targets, aggregates


def parse_post_data(post_data: str | None) -> dict[str, str]:
    return {
        key: value
        for key, value in urllib.parse.parse_qsl(post_data or "", keep_blank_values=True)
    }


def install_googleapis_compatibility(page) -> dict[str, int]:
    """Mirror jQuery and neutralize Google Fonts when googleapis is unreachable."""
    stats = {"jquery_mirrored": 0, "fonts_neutralized": 0}

    def jquery(route):
        response = route.fetch(url=JQUERY_MIRROR)
        route.fulfill(response=response)
        stats["jquery_mirrored"] += 1

    def fonts(route):
        route.fulfill(status=200, content_type="text/css", body="")
        stats["fonts_neutralized"] += 1

    page.route(re.compile(r"ajax\.googleapis\.com"), jquery)
    page.route(re.compile(r"fonts\.googleapis\.com"), fonts)
    return stats


def load_password() -> str:
    password = os.environ.get("DIGIPOS_PASSWORD")
    if password:
        return password
    here = Path(__file__).resolve()
    candidates = (Path(".env_encoded"), here.parents[3] / ".env_encoded")
    for candidate in candidates:
        if not candidate.is_file():
            continue
        for line in candidate.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("SECRET_DIGIPOS_PASSWORD="):
                encoded = line.split("=", 1)[1].strip()
                value = base64.b64decode(encoded + "===").decode("utf-8")
                if value:
                    return value
    raise RuntimeError("Set DIGIPOS_PASSWORD or provide .env_encoded")


def login(page, account: str, password: str) -> None:
    page.goto(f"{BASE_URL}/login", wait_until="domcontentloaded", timeout=90_000)
    page.wait_for_selector("#username", timeout=30_000)
    page.evaluate(
        "([username, password]) => { "
        "$('#username').val(username).trigger('change'); "
        "$('#password').val(password).trigger('change'); }",
        [account, password],
    )
    page.evaluate("$('button:contains(\"Sign in\")').trigger('click')")
    page.wait_for_url("**/home", timeout=30_000)


def open_monitoring(page) -> None:
    page.goto(MONITORING_URL, wait_until="domcontentloaded", timeout=90_000)
    page.wait_for_selector("#parentPhoneNo_pick", timeout=30_000)
    page.wait_for_function(
        "() => typeof $ !== 'undefined' && $.fn.DataTable.isDataTable('#dataInput')",
        timeout=60_000,
    )


def read_cells(page, account: str) -> tuple[list[CellTarget], list[dict]]:
    return discover_cells(account, page.evaluate(READ_CELL_OPTIONS))


def set_dates(page, report_date: str) -> None:
    day = datetime.strptime(report_date, "%Y-%m-%d")
    page.evaluate(
        SET_DATES,
        [
            day.strftime("%d %B %Y") + " 00:00:00",
            day.strftime("%d %B %Y") + " 23:59:59",
            report_date + " 00:00:00",
            report_date + " 23:59:59",
        ],
    )


def _status_code(payload: dict) -> str | None:
    value = payload.get("statusCode")
    return None if value is None else str(value)


def _raise_for_business_error(payload: dict) -> None:
    code = _status_code(payload)
    if code not in (None, "00"):
        description = str(payload.get("statusDesc") or "")
        raise RuntimeError(f"DigiPOS statusCode={code}: {description}")


def _footer_total(state: dict) -> int:
    text = state.get("info_text") or ""
    match = re.search(r"\bof\s+([\d,]+)\s+entries\b", text)
    if match:
        return int(match.group(1).replace(",", ""))
    total = state.get("total")
    if total is None:
        raise RuntimeError(f"could not read a stable table total: {text!r}")
    return int(total)


def _wait_table_settled(page, timeout_ms: int = 180_000) -> dict:
    page.wait_for_function(
        """() => {
            const table = $('#dataInput').DataTable();
            const text = ($('#dataInput_info').text() || '').trim();
            return !$('#loadingDialog').is(':visible') &&
                   !$('#modal-gagal').is(':visible') &&
                   (!table.processing || !table.processing()) &&
                   text && !text.includes('NaN') &&
                   !text.includes('Silakan pilih filter');
        }""",
        timeout=timeout_ms,
    )
    return page.evaluate(READ_TABLE_STATE)


def search_cell(page, target: CellTarget, report_date: str, timeout_ms: int = 180_000) -> dict:
    """Search in the observed working order: dates, cell, then real Search click."""
    set_dates(page, report_date)
    selection = page.evaluate(SELECT_CELL, target.parentcode)
    if not selection.get("ok") or selection.get("parentcode") != target.parentcode:
        raise RuntimeError(f"cell selection did not stick for {target.parentcode}")

    captured: list[dict[str, str]] = []

    def capture(route):
        captured.append(parse_post_data(route.request.post_data))
        route.continue_()

    page.route(DATA_ENDPOINT, capture)
    try:
        with page.expect_response(
            lambda response: DATA_ENDPOINT_PATH in response.url
            and response.request.method == "POST",
            timeout=timeout_ms,
        ) as response_info:
            page.click("#btn_search", timeout=30_000)
        response = response_info.value
    finally:
        page.unroute(DATA_ENDPOINT, capture)

    if response.status != 200:
        raise RuntimeError(f"search XHR HTTP {response.status}")
    payload = response.json()
    _raise_for_business_error(payload)
    if not captured:
        raise RuntimeError("Search response arrived without a captured request")
    request = captured[-1]
    expected_request = {
        "adNo": target.account,
        "parentcode": target.parentcode,
        "startDate": f"{report_date} 00:00:00",
        "endDate": f"{report_date} 23:59:59",
    }
    mismatches = {
        key: {"expected": value, "actual": request.get(key)}
        for key, value in expected_request.items()
        if request.get(key) != value
    }
    if mismatches:
        raise RuntimeError(f"search request did not match account/cell/date: {mismatches}")

    state = _wait_table_settled(page, timeout_ms)
    footer_total = _footer_total(state)
    rows = payload.get("data") or []
    records_filtered = payload.get("recordsFiltered")
    if records_filtered is not None and int(records_filtered) != footer_total:
        raise RuntimeError(
            f"search response total {records_filtered} does not match table footer {footer_total}"
        )
    if footer_total == 0 and rows:
        raise RuntimeError("search footer is empty but response returned rows")
    return {
        "footer_total": footer_total,
        "request": request,
        "search_response_rows": len(rows),
        "records_filtered": records_filtered,
        "status_code": _status_code(payload),
        "info_text": state["info_text"],
    }


def fetch_all_rows(
    page,
    target: CellTarget,
    report_date: str,
    footer_total: int,
    page_len: int = 5_000,
    timeout_ms: int = 180_000,
) -> tuple[list[dict], dict]:
    """Use the browser's own DataTables POST and enlarge only its page length."""
    if not 1 <= page_len <= MAX_PAGE_LENGTH:
        raise ValueError(f"page_len must be between 1 and {MAX_PAGE_LENGTH}")
    if footer_total == 0:
        return [], {"recordsFiltered": 0, "data": []}

    routed: list[dict] = []

    def rewrite_length(route):
        request = route.request
        post_data = request.post_data or ""
        content_type = request.headers.get("content-type", "")
        query = urllib.parse.parse_qsl(
            urllib.parse.urlsplit(request.url).query,
            keep_blank_values=True,
        )
        values = parse_post_data(post_data) if post_data else {}
        diagnostic = {
            "method": request.method,
            "content_type": content_type,
            "url_query_keys": sorted({key for key, _ in query}),
            "post_keys": sorted(values),
            "length_value": values.get("length"),
            "post_body_bytes": len(post_data.encode("utf-8")),
        }

        if "length" in values:
            body, replacements = re.subn(
                r"(^|&)length=-?\d+",
                lambda match: f"{match.group(1)}length={page_len}",
                post_data,
                count=1,
            )
            if replacements:
                values["length"] = str(page_len)
                routed.append({**values, **diagnostic})
                route.continue_(
                    post_data=body,
                    headers={**request.headers, "Content-Length": str(len(body.encode("utf-8")))},
                )
                return

        try:
            json_body = json.loads(post_data) if post_data else None
        except json.JSONDecodeError:
            json_body = None
        if isinstance(json_body, dict) and "length" in json_body:
            json_body["length"] = page_len
            body = json.dumps(json_body, separators=(",", ":"))
            values = {key: str(value) for key, value in json_body.items()}
            routed.append({**values, **diagnostic})
            route.continue_(
                post_data=body,
                headers={**request.headers, "Content-Length": str(len(body.encode("utf-8")))},
            )
            return

        query_values = dict(query)
        if "length" in query_values:
            query_values["length"] = str(page_len)
            parts = urllib.parse.urlsplit(request.url)
            url = urllib.parse.urlunsplit(
                parts._replace(query=urllib.parse.urlencode(query_values))
            )
            routed.append({**query_values, **diagnostic})
            route.continue_(url=url)
            return

        routed.append({**diagnostic, "_rewrite_error": "request has no length parameter"})
        route.continue_()

    page.route(DATA_ENDPOINT, rewrite_length)
    try:
        with page.expect_response(
            lambda response: DATA_ENDPOINT_PATH in response.url,
            timeout=timeout_ms,
        ) as response_info:
            page.evaluate("length => $('#dataInput').DataTable().page.len(length).draw()", page_len)
        response = response_info.value
    finally:
        page.unroute(DATA_ENDPOINT, rewrite_length)

    if response.status != 200:
        raise RuntimeError(f"full-row XHR HTTP {response.status}")
    payload = response.json()
    _raise_for_business_error(payload)
    if len(routed) != 1:
        raise RuntimeError(f"expected one rewritten XHR, intercepted {len(routed)}")
    request = routed[0]
    if request.get("_rewrite_error"):
        raise RuntimeError(f"{request['_rewrite_error']}: {request}")
    for key, value in {
        "length": str(page_len),
        "adNo": target.account,
        "parentcode": target.parentcode,
        "startDate": f"{report_date} 00:00:00",
        "endDate": f"{report_date} 23:59:59",
    }.items():
        if request.get(key) != value:
            raise RuntimeError(f"full-row XHR {key}={request.get(key)!r}; expected {value!r}")

    rows = payload.get("data") or []
    records_filtered = payload.get("recordsFiltered")
    if len(rows) != footer_total:
        raise RuntimeError(f"full XHR rows={len(rows)}; footer total={footer_total}")
    if records_filtered is not None and int(records_filtered) != footer_total:
        raise RuntimeError(f"recordsFiltered={records_filtered}; footer total={footer_total}")
    expected_date = datetime.strptime(report_date, "%Y-%m-%d").strftime("%d-%m-%Y")
    if any(str(row.get("tanggal") or "")[:10] != expected_date for row in rows):
        raise RuntimeError("full XHR contains a row outside the requested report date")
    return rows, {"recordsFiltered": records_filtered, "length": request["length"]}


def _to_int(value) -> int | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    return int(text) if re.fullmatch(r"-?\d+", text) else None


def row_to_cells(row: dict, index: int) -> list:
    transfer_id = str(row.get("wstransferid") or "").strip()
    numeric_id = _to_int(transfer_id)
    phone = str(row.get("sumber1") or "").strip()
    if phone.startswith("49"):
        phone = phone[2:]
    nomor_rs = (phone or "-") if row.get("transaction_type") == "Credit" else "-"
    return [
        index,
        str(row.get("tanggal") or "").strip(),
        numeric_id if numeric_id is not None else (transfer_id or "-"),
        _to_int(row.get("init_amount")) or 0,
        _to_int(row.get("amount_in")) or 0,
        _to_int(row.get("amount_out")) or 0,
        _to_int(row.get("last_amount")) or 0,
        str(row.get("transaction_type") or "").strip(),
        str(row.get("wstransfertype") or "").strip(),
        nomor_rs,
        str(row.get("wstransferdesc") or "").strip(),
    ]


def write_xlsx(path: Path, rows: list[list]) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.cell(row=1, column=1, value=TITLE)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for row_index in range(3, sheet.max_row + 1):
        for column in MONEY_COLUMNS:
            cell = sheet.cell(row=row_index, column=column)
            cell.number_format = "#,##0" if abs(cell.value or 0) >= 1000 else "0"
        id_cell = sheet.cell(row=row_index, column=3)
        if isinstance(id_cell.value, int):
            id_cell.number_format = "0"
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    workbook.save(path)


def artifact_filename(target: CellTarget, report_date: str) -> str:
    day = datetime.strptime(report_date, "%Y-%m-%d").strftime("%d-%m-%Y")
    return f"finpay-{target.account}-{target.cell_code}({day}to{day}).xlsx"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65_536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def row_count_from_xlsx(path: Path, report_date: str) -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [
            sheet.cell(row=2, column=column).value
            for column in range(1, len(HEADERS) + 1)
        ]
        if header != HEADERS:
            raise RuntimeError(f"workbook header mismatch: {path.name}")
        expected_date = datetime.strptime(report_date, "%Y-%m-%d").strftime("%d-%m-%Y")
        rows = 0
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if not any(value is not None for value in values):
                continue
            actual_date = str(values[1] or "").split(" ", 1)[0]
            if actual_date != expected_date:
                raise RuntimeError(
                    f"workbook date mismatch in {path.name}: "
                    f"expected {expected_date}, got {actual_date!r}"
                )
            rows += 1
        return rows
    finally:
        workbook.close()


def result_dict(target: CellTarget, report_date: str) -> dict:
    return {
        "account": target.account,
        "cluster_id": target.cluster_id,
        "cell_code": target.cell_code,
        "parentcode": target.parentcode,
        "cell_label": target.label,
        "report_date": report_date,
        "status": "pending",
        "artifact": None,
        "sha256": None,
        "footer_total": None,
        "validated_rows": None,
    }


def manifest_json(payload: dict) -> str:
    import json

    return json.dumps(payload, indent=2, sort_keys=True) + "\n"


__all__ = [
    "ACCOUNT_ROSTER",
    "MAX_PAGE_LENGTH",
    "CellTarget",
    "artifact_filename",
    "cluster_id_for",
    "discover_cells",
    "fetch_all_rows",
    "install_googleapis_compatibility",
    "load_password",
    "login",
    "manifest_json",
    "open_monitoring",
    "read_cells",
    "result_dict",
    "row_count_from_xlsx",
    "search_cell",
    "sha256_file",
    "write_xlsx",
    "row_to_cells",
    "set_dates",
]
