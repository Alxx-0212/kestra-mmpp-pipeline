#!/usr/bin/env python3
"""Scrape DigiPOS monitoring transactions into .xlsx via the DataTables XHR.

Use when the CMS Excel export button is missing or its download never fires.
Instead of walking table pages, the DataTables request is intercepted and its
page length is raised so one response carries every row; the JSON rows are then
written to Excel in the same shape as the real CMS export:

  row 1  title `CMS DigiPOS - Riwayat Saldo` (merged A1:K1)
  row 2  header  No / Transaction Date / Transaction ID / Saldo Awal / Kredit /
         Debet / Saldo Akhir / Transaction Type / Transaction / Nomor RS / Remarks
  row 3+ data, No and money columns numeric, `Nomor RS` text (never 8.23E+10)

Row count must equal the table footer `entries` total, else nothing is saved.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import download_digipos as dg  # noqa: E402

from openpyxl import Workbook  # noqa: E402

TITLE = "CMS DigiPOS - Riwayat Saldo"
HEADERS = ["No", "Transaction Date", "Transaction ID", "Saldo Awal", "Kredit", "Debet",
           "Saldo Akhir", "Transaction Type", "Transaction", "Nomor RS", "Remarks"]
MONEY_COLS = {3, 4, 5, 6}
ID_COLS = {2}
WIDTHS = {"A": 6.0, "B": 25.65, "C": 37.8, "D": 14.85, "E": 13.5, "F": 14.85,
          "G": 14.85, "H": 21.6, "I": 18.9, "J": 27.0, "K": 54.0}
ENDPOINT_GLOB = "**/deposit/monitoring-riwayat-data*"
ENDPOINT_PATH = "/deposit/monitoring-riwayat-data"


def dmy(date_str: str) -> str:
    return datetime.strptime(date_str, "%Y-%m-%d").strftime("%d-%m-%Y")


def money_format(value: int) -> str:
    return "#,##0" if abs(value) >= 1000 else "0"


def to_int(value) -> int | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not re.fullmatch(r"-?\d+", text):
        return None
    return int(text)


def strip49(value) -> str:
    """DigiPOS prefixes party numbers with 49; the export shows them without it."""
    text = "" if value is None else str(value).strip()
    return text[2:] if text.startswith("49") else text


def nomor_rs(row: dict) -> str:
    """Verified against the real 28-09 export: counterparty only on Credit rows."""
    if (row.get("transaction_type") or "").strip() != "Credit":
        return "-"
    return strip49(row.get("sumber1")) or "-"


def transfer_id(row: dict):
    """Digit-only IDs are exported as numbers; prefixed IDs stay text."""
    text = str(row.get("wstransferid") or "").strip()
    number = to_int(text)
    return number if number is not None else (text or "-")


def row_to_cells(row: dict, index: int) -> list:
    return [
        index,
        str(row.get("tanggal") or "").strip(),
        transfer_id(row),
        to_int(row.get("init_amount")) or 0,
        to_int(row.get("amount_in")) or 0,
        to_int(row.get("amount_out")) or 0,
        to_int(row.get("last_amount")) or 0,
        str(row.get("transaction_type") or "").strip(),
        str(row.get("wstransfertype") or "").strip(),
        nomor_rs(row),
        str(row.get("wstransferdesc") or "").strip(),
    ]


def write_xlsx(path: str, rows: list[list]) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value=TITLE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.append(HEADERS)
    for cells in rows:
        ws.append(cells)
    for r in range(3, ws.max_row + 1):
        for c in MONEY_COLS:
            ws.cell(row=r, column=c + 1).number_format = money_format(ws.cell(row=r, column=c + 1).value)
        for c in ID_COLS:
            cell = ws.cell(row=r, column=c + 1)
            if isinstance(cell.value, int):
                cell.number_format = "0"
    for letter, width in WIDTHS.items():
        ws.column_dimensions[letter].width = width
    wb.save(path)


def fetch_all(page, date_str: str, page_len: int) -> tuple[list[list], int]:
    info = dg.set_date_and_search(page, date_str)
    expected = int(info.get("totalEntries") or 0)
    if expected <= 0:
        raise RuntimeError(f"{date_str}: table shows 0 entries")

    def route_it(route):
        req = route.request
        if ENDPOINT_PATH in req.url and req.method == "POST":
            data = re.sub(r"(^|&)length=\d+", lambda m: f"{m.group(1)}length={page_len}",
                          req.post_data or "", count=1)
            route.continue_(post_data=data, headers={**req.headers, "Content-Length": str(len(data))})
        else:
            route.continue_()

    page.route(ENDPOINT_GLOB, route_it)
    with page.expect_response(lambda r: ENDPOINT_PATH in r.url and r.request.method == "POST",
                              timeout=180000) as resp:
        page.evaluate("n => $('#dataInput').DataTable().page.len(n).draw();", page_len)
    page.unroute(ENDPOINT_GLOB)
    if resp.value.status != 200:
        raise RuntimeError(f"{date_str}: XHR HTTP {resp.value.status}")
    payload = resp.value.json()
    data = payload.get("data") or []
    if not data:
        raise RuntimeError(f"{date_str}: XHR returned no rows")
    if len(data) != expected:
        raise RuntimeError(f"{date_str}: footer total={expected}, XHR rows={len(data)}")
    rows = [row_to_cells(r, i) for i, r in enumerate(data, 1)]
    if any(len(c) != len(HEADERS) for c in rows):
        raise RuntimeError(f"{date_str}: column count mismatch")
    return rows, expected


def scrape_one(page, username: str, date_str: str, out_dir: str, page_len: int):
    cid = dg.cluster_id(username)
    d = dmy(date_str)
    path = str(Path(out_dir) / cid / f"finpay-{cid}({d}to{d}).xlsx")
    rows, expected = fetch_all(page, date_str, page_len)
    write_xlsx(path, rows)
    size = Path(path).stat().st_size
    print(f"  OK {Path(path).name}: rows={len(rows):,} table_total={expected:,} bytes={size:,}", flush=True)
    return {"cluster": cid, "user": username, "date": date_str, "file": path,
            "rows": len(rows), "table_total": expected, "bytes": size}


def main() -> int:
    ap = argparse.ArgumentParser(description="Scrape DigiPOS monitoring rows via the DataTables XHR")
    ap.add_argument("--users", nargs="+", default=dg.DEFAULT_USERS)
    ap.add_argument("--password", default=dg.PASSWORD)
    ap.add_argument("--date", required=True, help="YYYY-MM-DD")
    ap.add_argument("--output-dir", required=True)
    ap.add_argument("--page-len", type=int, default=5000, help="rows per XHR (one request per cluster)")
    ap.add_argument("--max-attempts", type=int, default=3)
    ap.add_argument("--results", default=None)
    args = ap.parse_args()

    from playwright.sync_api import sync_playwright

    report = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            for user in args.users:
                last = None
                for attempt in range(1, args.max_attempts + 1):
                    page = None
                    try:
                        print(f"{user} / {args.date} attempt {attempt}/{args.max_attempts}", flush=True)
                        page = browser.new_page(viewport={"width": 1440, "height": 1000})
                        dg.login(page, user, args.password)
                        report.append(scrape_one(page, user, args.date, args.output_dir, args.page_len))
                        last = None
                        break
                    except Exception as e:  # noqa: BLE001
                        last = e
                        print(f"  attempt {attempt} failed: {e}", flush=True)
                        time.sleep(5 * attempt)
                    finally:
                        if page:
                            try:
                                page.close()
                            except Exception:
                                pass
                if last is not None:
                    report.append({"cluster": dg.cluster_id(user), "user": user, "date": args.date, "error": str(last)})
        finally:
            browser.close()

    if args.results:
        Path(args.results).parent.mkdir(parents=True, exist_ok=True)
        Path(args.results).write_text(json.dumps(report, indent=2))
    ok = sum(1 for r in report if "error" not in r)
    print(json.dumps({"ok": ok, "failed": len(report) - ok, "date": args.date, "output_dir": args.output_dir}, indent=2), flush=True)
    return 1 if ok != len(report) else 0


if __name__ == "__main__":
    raise SystemExit(main())
