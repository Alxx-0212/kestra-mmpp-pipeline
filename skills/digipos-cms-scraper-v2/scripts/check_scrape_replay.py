"""Self-check: replay cached XHR rows through the writer and diff vs real exports.

Covers all six clusters for 28-09-2026, where real CMS exports exist on disk.
Requires a scrape first so the cached raw JSON exists.
"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0, "/home/mmpp/.hermes/profiles/alex/skills/digipos-cms-scraper/scripts")
import scrape_digipos_table as s

REAL = Path("/home/mmpp/projects/kestra-mmpp-pipeline/tmp/kestra-wd/finpay-inbox")
OUT = Path("/tmp/format_probe/replay")
DATE = "2026-09-28"
DMY = "28-09-2026"

total_bad = 0
for cache in sorted(Path("/tmp/format_probe").glob(f"raw_{DATE}_*_A.json")):
    user = cache.stem.replace(f"raw_{DATE}_", "")
    cid = user.replace("_A", "")
    payload = json.loads(cache.read_text())
    rows = payload["rows"]
    assert len(rows) == payload["table_total"], (cid, len(rows), payload["table_total"])
    cells = [s.row_to_cells(r, i) for i, r in enumerate(rows, 1)]
    out = OUT / cid / f"finpay-{cid}({DMY}to{DMY}).xlsx"
    s.write_xlsx(str(out), cells)

    real = REAL / cid / f"finpay-{cid}({DMY}to{DMY}).xlsx"
    if not real.exists():
        print(f"{cid}: no real export to compare, skipped")
        continue
    a = load_workbook(real, data_only=True).active
    b = load_workbook(out, data_only=True).active
    if (a.max_row, a.max_column) != (b.max_row, b.max_column):
        print(f"{cid}: SHAPE MISMATCH real={a.max_row}x{a.max_column} replay={b.max_row}x{b.max_column}")
        total_bad += 1
        continue
    bad = []
    for r in range(1, a.max_row + 1):
        for c in range(1, a.max_column + 1):
            va, vb = a.cell(r, c).value, b.cell(r, c).value
            if va != vb or type(va) is not type(vb):
                bad.append((r, c, a.cell(2, c).value, va, vb))
    total_bad += len(bad)
    status = "OK" if not bad else "DIFF"
    print(f"{cid}: {status} rows={a.max_row-2} diffs={len(bad)}")
    for d in bad[:6]:
        print(f"    r{d[0]} c{d[1]} {d[2]}: real={d[3]!r}({type(d[3]).__name__}) replay={d[4]!r}({type(d[4]).__name__})")

print("TOTAL_DIFFS", total_bad)
assert total_bad == 0, "replay does not reproduce the real export"
print("REPLAY_MATCH_OK")
