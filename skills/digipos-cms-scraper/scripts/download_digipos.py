#!/usr/bin/env python3
"""
DigiPOS CMS Scraper — browser automation via Playwright.

Hard guarantees:
- every requested date must have at least 1 transaction
- total entries shown by DataTables must equal downloaded CSV data rows
- CSV dates must match requested date
- retry with fresh login/page when web/export is flaky
"""

import argparse
import csv
import os
import re
import sys
import time
from datetime import datetime, timedelta

DEFAULT_USERS = ["411311_A", "421306_A", "421307_A", "421315_A", "421318_A", "421320_A"]
PASSWORD = "TselDigipos123"
BASE_URL = "https://digipos-cms.finpay.id"
OUTPUT_DIR = "/home/mmpp/projects/kestra-mmpp-pipeline/tmp/kestra-wd/finpay-inbox"


def cluster_id(username: str) -> str:
    return username.replace("_A", "")


def target_path_for(username: str, date_str: str, output_dir: str) -> str:
    cid = cluster_id(username)
    d1 = datetime.strptime(date_str, "%Y-%m-%d").strftime("%d-%m-%Y")
    return os.path.join(output_dir, cid, f"finpay-{cid}({d1}to{d1}).csv")


def target_xlsx_path_for(username: str, date_str: str, output_dir: str) -> str:
    cid = cluster_id(username)
    d1 = datetime.strptime(date_str, "%Y-%m-%d").strftime("%d-%m-%Y")
    return os.path.join(output_dir, cid, f"finpay-{cid}({d1}to{d1}).xlsx")


def safe_remove(path: str) -> None:
    if os.path.exists(path):
        os.remove(path)
        print(f"  Removed stale file: {path}", flush=True)


def read_csv_data_rows(path: str):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    return rows[1:] if rows else []


def read_xlsx_rows(path: str):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    wb.close()
    return rows


def normalize_export_rows(rows):
    # Excel export may prepend title/filter rows before the table header.
    # Keep only the actual table starting at header row whose first cells are No / Transaction Date.
    for i, row in enumerate(rows):
        cells = [str(c).strip() for c in row]
        if len(cells) >= 2 and cells[0] == "No" and cells[1] == "Transaction Date":
            return rows[i:]
    return rows


def write_csv_from_rows(path: str, rows) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def login(page, username: str, password: str) -> None:
    page.goto(f"{BASE_URL}/login", wait_until="networkidle")
    page.evaluate("([u, p]) => { $('#username').val(u).trigger('change'); $('#password').val(p).trigger('change'); }", [username, password])
    page.evaluate("$('button:contains(\"Sign in\")').trigger('click')")
    page.wait_for_url("**/home", timeout=20000)
    print(f"  Logged in as {username}", flush=True)


def set_date_and_search(page, date_str: str) -> dict:
    page.goto(f"{BASE_URL}/deposit/monitoring-riwayat", wait_until="networkidle")
    page.wait_for_selector("input[name=date1]", timeout=15000)

    dt = datetime.strptime(date_str, "%Y-%m-%d")
    display_start = dt.strftime("%d %B %Y") + " 00:00:00"
    display_end = dt.strftime("%d %B %Y") + " 23:59:59"
    api_start = dt.strftime("%Y-%m-%d") + " 00:00:00"
    api_end = dt.strftime("%Y-%m-%d") + " 23:59:59"

    page.evaluate("""([startVal, endVal, apiStart, apiEnd]) => {
        const setValue = (sel, val) => {
            const el = document.querySelector(sel);
            if (!el) return;
            const setter = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value').set;
            setter.call(el, val);
            el.dispatchEvent(new Event('input', { bubbles: true }));
            el.dispatchEvent(new Event('change', { bubbles: true }));
        };
        setValue('input[name="date1"]', startVal);
        setValue('input[name="date2"]', endVal);
        setValue('#dateStart', apiStart);
        setValue('#dateEnd', apiEnd);
    }""", [display_start, display_end, api_start, api_end])

    # Use real browser click; jQuery trigger can leave table in no-filter state.
    page.click('#btn_search')
    page.wait_for_function("() => typeof $ !== 'undefined' && $.fn.DataTable.isDataTable('#dataInput')", timeout=15000)
    # DigiPOS can be very slow. DataTables processing() may become false before
    # the site loading overlay disappears or before footer text updates.
    page.wait_for_function("""() => {
        const loading = $('#loadingDialog').is(':visible');
        const modal = $('#modal-gagal').is(':visible');
        const info = $('#dataInput_info').text().trim();
        return !loading && !modal && info && !info.includes('Silakan pilih filter') && !info.includes('NaN');
    }""", timeout=180000)
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)

    info = page.evaluate("""() => {
        const dt = $('#dataInput').DataTable();
        const info = dt.page.info();
        const infoText = $('#dataInput_info').text().trim();
        const pageRows = $('#dataInput tbody tr').filter(function(){
            return !$(this).find('td.dataTables_empty').length;
        }).length;
        return {
            recordsTotal: info.recordsTotal || 0,
            recordsDisplay: info.recordsDisplay ?? info.recordsFiltered ?? 0,
            infoText,
            pageRows,
            firstText: $('#dataInput tbody tr:first').text().trim()
        };
    }""")

    # DataTables info text is source of truth shown under table:
    # "Showing 1 to 10 of 3,111 entries" or "Showing 0 to 0 of 0 entries"
    m = re.search(r"of\s+([\d,]+)\s+entries", info.get("infoText") or "")
    shown_total = int(m.group(1).replace(",", "")) if m else int(info.get("recordsDisplay") or 0)
    api_total = int(info.get("recordsDisplay") or info.get("recordsTotal") or 0)
    total_entries = shown_total or api_total
    info["totalEntries"] = total_entries
    return info


def export_and_validate(page, username: str, date_str: str, output_dir: str, csv_copy: bool = True) -> str:
    path = target_path_for(username, date_str, output_dir)
    xlsx_path = target_xlsx_path_for(username, date_str, output_dir)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    safe_remove(path)
    safe_remove(xlsx_path)

    page.evaluate("$('#dataInput').DataTable().page.len(-1).draw()")
    page.wait_for_function("() => { const t = $('#dataInput').DataTable(); return !t.processing || !t.processing(); }", timeout=30000)
    page.wait_for_load_state("networkidle")
    time.sleep(2)

    # Use Excel export, not CSV export. CSV export can serialize Nomor RS
    # phone-like values as scientific notation (for example 8E+10).
    with page.expect_download(timeout=45000) as download_info:
        page.evaluate("$('.buttons-excel').trigger('click')")
    download_info.value.save_as(xlsx_path)

    rows = normalize_export_rows(read_xlsx_rows(xlsx_path))
    if csv_copy:
        write_csv_from_rows(path, rows)

    data_rows = rows[1:] if rows else []
    target_ddmmyyyy = datetime.strptime(date_str, "%Y-%m-%d").strftime("%d-%m-%Y")
    unique_dates = sorted({r[1][:10] for r in data_rows if len(r) > 1 and r[1]})
    nomor_rs_values = [r[9].strip() for r in data_rows if len(r) > 9 and r[9].strip() and r[9].strip() != "-"]
    scientific_nomor_rs = [v for v in nomor_rs_values if re.search(r"\d+\.?\d*E\+?\d+", v, re.I)]

    if not data_rows:
        safe_remove(path)
        safe_remove(xlsx_path)
        raise RuntimeError("Excel export has 0 data rows")
    if unique_dates != [target_ddmmyyyy]:
        safe_remove(path)
        safe_remove(xlsx_path)
        raise RuntimeError(f"Excel date mismatch: expected {target_ddmmyyyy}, got {unique_dates[:5]}")
    if scientific_nomor_rs:
        safe_remove(path)
        safe_remove(xlsx_path)
        raise RuntimeError(f"Nomor RS still scientific notation after Excel export: {scientific_nomor_rs[:5]}")

    print(f"  Saved Excel: {os.path.basename(xlsx_path)} ({os.path.getsize(xlsx_path):,} B, rows={len(data_rows):,})", flush=True)
    if csv_copy:
        print(f"  Saved CSV copy: {os.path.basename(path)} ({os.path.getsize(path):,} B, rows={len(data_rows):,})", flush=True)
        return path
    return xlsx_path


def download_for_user(browser, username: str, date_str: str, output_dir: str, password: str, max_attempts: int, csv_copy: bool = True) -> str:
    last_error = None
    for attempt in range(1, max_attempts + 1):
        page = None
        try:
            print(f"  Attempt {attempt}/{max_attempts}", flush=True)
            page = browser.new_page(accept_downloads=True, viewport={"width": 1280, "height": 800})
            login(page, username, password)

            # Search can be flaky; retry inside same fresh login with refresh.
            info = None
            for search_try in range(1, 4):
                info = set_date_and_search(page, date_str)
                print(
                    f"  Search {search_try}/3: info='{info.get('infoText')}', totalEntries={info.get('totalEntries')}, pageRows={info.get('pageRows')}",
                    flush=True,
                )
                if int(info.get("totalEntries") or 0) > 0:
                    break
                page.reload(wait_until="networkidle")
                time.sleep(2)

            total_entries = int((info or {}).get("totalEntries") or 0)
            if total_entries <= 0:
                raise RuntimeError(f"Table shows 0 entries for {username} {date_str}; expected at least one transaction")
            path = export_and_validate(page, username, date_str, output_dir, csv_copy=csv_copy)
            xlsx_rows = normalize_export_rows(read_xlsx_rows(target_xlsx_path_for(username, date_str, output_dir)))
            data_rows = xlsx_rows[1:] if xlsx_rows else []
            if len(data_rows) != total_entries:
                safe_remove(target_path_for(username, date_str, output_dir))
                safe_remove(target_xlsx_path_for(username, date_str, output_dir))
                raise RuntimeError(f"Excel row mismatch: table total entries={total_entries}, exported data rows={len(data_rows)}")
            return path
        except Exception as e:
            last_error = e
            print(f"  Attempt {attempt} failed: {e}", flush=True)
            safe_remove(target_path_for(username, date_str, output_dir))
            safe_remove(target_xlsx_path_for(username, date_str, output_dir))
            time.sleep(3 * attempt)
        finally:
            if page:
                try:
                    page.close()
                except Exception:
                    pass
    raise RuntimeError(f"All attempts failed for {username} {date_str}: {last_error}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Download DigiPOS CMS data via browser automation")
    parser.add_argument("--users", nargs="+", default=DEFAULT_USERS)
    parser.add_argument("--password", default=PASSWORD)
    parser.add_argument("--date", help="Single date (YYYY-MM-DD), defaults to yesterday")
    parser.add_argument("--start", help="Start date (YYYY-MM-DD)")
    parser.add_argument("--end", help="End date (YYYY-MM-DD)")
    parser.add_argument("--output-dir", default=OUTPUT_DIR)
    parser.add_argument("--max-attempts", type=int, default=5)
    parser.add_argument("--headless", action="store_true", default=True)
    parser.add_argument("--no-headless", action="store_true")
    parser.add_argument("--excel-only", action="store_true", help="Save only .xlsx files; do not write CSV copies")
    args = parser.parse_args()

    if args.date:
        dates = [args.date]
    elif args.start and args.end:
        start = datetime.strptime(args.start, "%Y-%m-%d")
        end = datetime.strptime(args.end, "%Y-%m-%d")
        dates = [(start + timedelta(days=i)).strftime("%Y-%m-%d") for i in range((end - start).days + 1)]
    else:
        dates = [(datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")]

    os.makedirs(args.output_dir, exist_ok=True)
    headless = not args.no_headless
    print(f"Dates: {dates[0]} -> {dates[-1]}")
    print(f"Users: {len(args.users)}")
    print(f"Headless: {headless}")
    print(f"Output: {args.output_dir}\n")

    from playwright.sync_api import sync_playwright

    ok = fail = 0
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless)
        try:
            for user in args.users:
                for date_str in dates:
                    print(f"{user} / {date_str}...", flush=True)
                    try:
                        download_for_user(browser, user, date_str, args.output_dir, args.password, args.max_attempts, csv_copy=not args.excel_only)
                        ok += 1
                    except Exception as e:
                        print(f"  FAIL: {e}", flush=True)
                        fail += 1
        finally:
            browser.close()

    print(f"\nDone: {ok} OK, {fail} failed")
    return 1 if fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
