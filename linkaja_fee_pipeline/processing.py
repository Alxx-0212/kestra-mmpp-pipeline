"""LinkAja fee extraction and aggregation."""

from __future__ import annotations

import csv
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


REQUIRED_COLUMNS = ("Finalized Date", "Transaction Scenario", "Debit", "Credit", "Fee")
EXPECTED_SCENARIO = "Digipos B2B Transfer"
IN_CLUSTER_SCENARIO = "Digipos B2B Transfer In Cluster"
EXPECTED_FEE_PER_ROW = Decimal("200")

REPORT_DATE_HEADER = "REPORT DATE"
CLUSTER_ID_HEADER = "CLUSTER ID"
EXPECTED_FEE_HEADER = "EXPECTED RECHARGE OUT CLUSTER FEE"
IN_CLUSTER_FEE_HEADER = "DIGIPOS B2B TRANSFER IN CLUSTER FEE"
TOTAL_FEE_HEADER = "TOTAL FEE"
SOURCE_FILE_HEADER = "SOURCE FILE"
SOURCE_ROWS_HEADER = "SOURCE ROWS"

SHEET_HEADERS = [
    REPORT_DATE_HEADER,
    CLUSTER_ID_HEADER,
    EXPECTED_FEE_HEADER,
    IN_CLUSTER_FEE_HEADER,
    TOTAL_FEE_HEADER,
    SOURCE_FILE_HEADER,
    SOURCE_ROWS_HEADER,
]


@dataclass(frozen=True)
class LinkAjaFeeSummary:
    cluster_id: str
    source_file: str
    source_rows: int
    report_dates: list[str]
    rows: list[dict[str, Any]]

    def to_json_dict(self) -> dict[str, Any]:
        return {
            "cluster_id": self.cluster_id,
            "source_file": self.source_file,
            "source_rows": self.source_rows,
            "report_dates": self.report_dates,
            "rows": self.rows,
        }


def display_name(value: str | Path) -> str:
    text = str(value).rstrip("/")
    return text.rsplit("/", 1)[-1] or text


def cluster_id_from_filename(value: str | Path) -> str | None:
    match = re.search(r"laporan-(\d+)-", display_name(value))
    return match.group(1) if match else None


def cluster_id_from_rows(rows: Iterable[dict[str, Any]]) -> str | None:
    for row in rows:
        for column in ("Top Organization", "Parent Organization", "Organization"):
            value = str(row.get(column, "") or "")
            match = re.match(r"\s*(\d+)-", value)
            if match:
                return match.group(1)
    return None


def parse_amount(value: Any, source_name: str, row_number: int, column: str) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if math.isnan(value):
            return Decimal("0")
        return Decimal(str(value))

    text = str(value).strip()
    if not text or text.lower() == "nan":
        return Decimal("0")

    normalized = text.replace(",", "")
    if normalized.startswith("(") and normalized.endswith(")"):
        normalized = "-" + normalized[1:-1]

    try:
        return Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(
            f"{source_name}:{row_number}: cannot parse {column} value {value!r} as a number"
        ) from exc


def normalize_report_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")

    text = str(value).strip()
    if not text:
        return ""

    for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, date_format).strftime("%d/%m/%Y")
        except ValueError:
            pass

    try:
        return datetime.fromisoformat(text).strftime("%d/%m/%Y")
    except ValueError:
        return text


def date_sort_key(value: str) -> tuple[int, Any]:
    for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return (0, datetime.strptime(value, date_format).date())
        except ValueError:
            pass
    return (1, value)


def decimal_to_sheet_value(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def load_linkaja_rows(
    source_path: str | Path,
    filename_hint: str | None = None,
) -> list[dict[str, Any]]:
    path = Path(source_path)
    suffix = path.suffix.lower() or Path(filename_hint or "").suffix.lower()

    if suffix == ".csv" or suffix == "":
        return _load_csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx_rows(path)

    raise ValueError(f"Unsupported LinkAja input file extension: {suffix or '<none>'}")


def _load_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        _validate_columns(reader.fieldnames or [], display_name(path))
        return [dict(row) for row in reader]


def _load_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read LinkAja .xlsx files") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    all_rows: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        row_values = worksheet.iter_rows(values_only=True)
        headers = None
        for raw_header in row_values:
            if raw_header and any(cell is not None and str(cell).strip() for cell in raw_header):
                headers = [str(cell).strip() if cell is not None else "" for cell in raw_header]
                break

        if not headers:
            continue

        _validate_columns(headers, f"{display_name(path)}:{worksheet.title}")
        for raw_row in row_values:
            if not raw_row or not any(cell is not None and str(cell).strip() for cell in raw_row):
                continue
            all_rows.append({
                header: raw_row[index] if index < len(raw_row) else ""
                for index, header in enumerate(headers)
            })

    if not all_rows:
        raise ValueError(f"{display_name(path)}: no data rows found")

    return all_rows


def _validate_columns(fieldnames: list[str], source_name: str) -> None:
    missing = [column for column in REQUIRED_COLUMNS if column not in fieldnames]
    if missing:
        missing_text = ", ".join(missing)
        raise ValueError(f"{source_name}: missing required column(s): {missing_text}")


def aggregate_linkaja_fee_file(
    source_path: str | Path,
    filename_hint: str | None = None,
) -> LinkAjaFeeSummary:
    rows = load_linkaja_rows(source_path, filename_hint=filename_hint)
    source_name = display_name(filename_hint or source_path)
    cluster_id = cluster_id_from_filename(filename_hint or source_path) or cluster_id_from_rows(rows)

    if not cluster_id:
        raise ValueError(
            f"{source_name}: cannot determine cluster ID from filename or organization columns"
        )

    expected_counts: dict[str, int] = defaultdict(int)
    in_cluster_fees: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    source_rows_by_date: dict[str, int] = defaultdict(int)

    for row_index, row in enumerate(rows, start=2):
        report_date = normalize_report_date(row.get("Finalized Date"))
        if not report_date:
            raise ValueError(f"{source_name}:{row_index}: Finalized Date is blank")

        parse_amount(row.get("Debit"), source_name, row_index, "Debit")
        credit = parse_amount(row.get("Credit"), source_name, row_index, "Credit")
        fee = parse_amount(row.get("Fee"), source_name, row_index, "Fee")
        scenario = str(row.get("Transaction Scenario", "") or "").strip()

        source_rows_by_date[report_date] += 1
        if scenario == EXPECTED_SCENARIO and credit != 0:
            expected_counts[report_date] += 1
        if scenario == IN_CLUSTER_SCENARIO:
            in_cluster_fees[report_date] += fee

    report_dates = sorted(source_rows_by_date, key=date_sort_key)
    output_rows = []
    for report_date in report_dates:
        expected_fee = Decimal(expected_counts.get(report_date, 0)) * EXPECTED_FEE_PER_ROW
        in_cluster_fee = in_cluster_fees.get(report_date, Decimal("0"))
        total_fee = expected_fee + in_cluster_fee
        output_rows.append({
            REPORT_DATE_HEADER: report_date,
            CLUSTER_ID_HEADER: cluster_id,
            EXPECTED_FEE_HEADER: decimal_to_sheet_value(expected_fee),
            IN_CLUSTER_FEE_HEADER: decimal_to_sheet_value(in_cluster_fee),
            TOTAL_FEE_HEADER: decimal_to_sheet_value(total_fee),
            SOURCE_FILE_HEADER: source_name,
            SOURCE_ROWS_HEADER: source_rows_by_date[report_date],
        })

    return LinkAjaFeeSummary(
        cluster_id=cluster_id,
        source_file=source_name,
        source_rows=len(rows),
        report_dates=report_dates,
        rows=output_rows,
    )


def write_linkaja_fee_json(summary: LinkAjaFeeSummary, output_path: str | Path) -> None:
    Path(output_path).write_text(
        json.dumps(summary.to_json_dict(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_linkaja_fee_json(input_path: str | Path) -> dict[str, Any]:
    return json.loads(Path(input_path).read_text(encoding="utf-8"))
