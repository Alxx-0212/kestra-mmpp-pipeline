"""Offline design/formula-request tests; no Excel/Google formula engine is used."""
import copy
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from finpay_pipeline.cell_design import dummy_cluster_report
from finpay_pipeline.cell_calculation import build_cluster_report
from finpay_pipeline.tests import test_cell_calculation as fixtures
from finpay_pipeline.cell_sheets import (
    column_letter, mandiri_reset_dates, render_cell_matrix as current_matrix, write_cell_design_workbook,
)


def render_cell_matrix(*args, **kwargs):
    """Keep the published wide-v2 regression contract explicit."""
    return current_matrix(*args, **kwargs, presentation_version=2)


class CellSheetDesignTests(unittest.TestCase):
    def setUp(self):
        self.report = dummy_cluster_report("412110_B", ("412111", "412112"))

    def input(self, matrix, day, code):
        return next(item for item in matrix["editable_inputs"]
                    if item["report_date"] == day and item["cell_code"] == code)

    def row(self, matrix, day, key):
        start = next(item["start_row"] for item in matrix["date_ranges"] if item["report_date"] == day)
        return next(index + 1 for index, value in enumerate(matrix["row_keys"])
                    if index + 1 >= start and value == key)

    def test_two_cell_layout_is_twenty_columns_with_identical_date_rows(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        self.assertEqual(len(matrix["values"][0]), 20)
        self.assertEqual([(block["start_column"], block["end_column"]) for block in matrix["blocks"]],
                         [(1, 6), (8, 13), (15, 20)])
        self.assertIn("412111", matrix["values"][1][7])
        self.assertIn("412112", matrix["values"][1][14])
        self.assertEqual(matrix["values"][0][0], "REPORT DATE")
        self.assertEqual(matrix["layout_version"], 2)
        self.assertNotIn("PUBLICATION", [value for row in matrix["values"] for value in row])
        for row in matrix["values"][2:]:
            self.assertEqual(row[0], row[7])
            self.assertEqual(row[0], row[14])
            self.assertEqual(row[2], row[9])
            self.assertEqual(row[2], row[16])
            self.assertEqual(row[6], "")
            self.assertEqual(row[13], "")

    def test_three_cell_layout_and_column_names(self):
        report = dummy_cluster_report("422040_B", ("422041", "422042", "422043"))
        matrix = render_cell_matrix(report, dummy=True)
        self.assertEqual(len(matrix["values"][0]), 27)
        self.assertEqual(matrix["blocks"][-1]["end_column"], 27)
        self.assertEqual([column_letter(number) for number in (11, 18, 25, 27)], ["K", "R", "Y", "AA"])

    def test_all_numeric_detail_saldos_equal_sum_of_cells(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        for key, row in zip(matrix["row_keys"], matrix["numeric_preview"]):
            if type(row[5]) is int:
                self.assertEqual(row[5], row[12] + row[19], key)
        self.assertEqual(matrix["values"][2][5], "=SUM(M3,T3)")

    def test_formulas_reference_correct_cell_columns_not_legacy_d(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        for code, col in (("412111", "K"), ("412112", "R")):
            bank = self.input(matrix, "2026-10-13", code)
            formula = matrix["values"][bank["row"] - 1][bank["column"] - 1]
            self.assertTrue(formula.startswith("=" + col))
            self.assertNotIn("D", formula)
            cell_saldo = 13 if code == "412111" else 20
            detail = self.row(matrix, "2026-10-13", "DETAIL:SELLTHRU")
            saldo_formula = matrix["values"][detail - 1][cell_saldo - 1]
            self.assertIn(col + str(detail), saldo_formula)

    def test_manual_reset_only_requested_and_preceding_dates(self):
        reset = mandiri_reset_dates(self.report, ["2026-10-15"])
        self.assertEqual(reset, {"2026-10-14", "2026-10-15"})
        matrix = render_cell_matrix(
            self.report, dummy=True, reset_dates=reset,
            manual_values={("2026-10-13", "412111"): 45000,
                           ("2026-10-14", "412111"): 33333},
        )
        old, refreshed = self.input(matrix, "2026-10-13", "412111"), self.input(matrix, "2026-10-14", "412111")
        self.assertEqual(matrix["values"][old["row"] - 1][old["column"] - 1], 45000)
        self.assertTrue(matrix["values"][refreshed["row"] - 1][refreshed["column"] - 1].startswith("=K"))
        self.assertEqual(len(matrix["manual_reset_audit"]), 1)
        self.assertEqual(matrix["manual_reset_audit"][0]["previous_value"], 33333)

    def test_latest_day_bank_default_is_pending_not_zero(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        bank = self.input(matrix, "2026-10-15", "412111")
        row = matrix["values"][bank["row"] - 1]
        self.assertEqual(row[bank["column"] - 1], "PENDING NEXT DAY")
        self.assertIn("PENDING NEXT DAY", row[3])
        self.assertEqual(matrix["numeric_preview"][bank["row"] - 1][3], "PENDING NEXT DAY")

    def test_no_invoice_linkaja_or_legacy_panels(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        text = str(matrix["values"]).upper()
        for forbidden in ("INVOICE", "LINKAJA", "CASH-IN", "CASH-OUT"):
            self.assertNotIn(forbidden, text)

    def test_combined_negative_mandiri_is_split_after_netting_cells(self):
        fixture = fixtures.CellCalculationTests()
        fixture.setUp()
        one = fixture.prepare("412111", [fixture.row(amount=0, debit=100)])
        two = fixture.prepare("412112", [fixture.row(amount=20)])
        report = build_cluster_report(fixture.config, [one, two], "2026-10-13")
        matrix = render_cell_matrix(report, dummy=True)
        row = matrix["row_keys"].index("MANDIRI:net:ST")
        self.assertEqual(matrix["numeric_preview"][row][3:5], [0, 80])
        self.assertEqual(matrix["numeric_preview"][row][10], -100)
        self.assertEqual(matrix["numeric_preview"][row][17], 20)
        self.assertEqual(matrix["values"][row][3], f"=MAX(SUM(K{row + 1},R{row + 1}),0)")
        self.assertEqual(matrix["values"][row][4], f"=MAX(-SUM(K{row + 1},R{row + 1}),0)")

    def test_numeric_preview_cash_difference_sign_matches_finance_layout(self):
        matrix = render_cell_matrix(self.report, dummy=True,
                                    manual_values={("2026-10-13", "412111"): 45000})
        row = self.row(matrix, "2026-10-13", "CASH:cash:SELISIH")
        self.assertEqual(matrix["numeric_preview"][row - 1][10], -55000)
        self.assertEqual(matrix["numeric_preview"][row - 1][11], "KURANG BAYAR")

    def test_no_category_day_is_zero_movement_and_carries_prior_saldo(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        last = self.row(matrix, "2026-10-13", "DETAIL:SELLTHRUFEE")
        empty = self.row(matrix, "2026-10-14", "DETAIL:SELLTHRUFEE")
        self.assertEqual(matrix["numeric_preview"][empty - 1][19], matrix["numeric_preview"][last - 1][19])

    def test_bad_coverage_manual_or_float_amount_cannot_render(self):
        changed = copy.deepcopy(self.report)
        changed["days"].pop(1)
        with self.assertRaises(ValueError):
            render_cell_matrix(changed)
        with self.assertRaises(ValueError):
            render_cell_matrix(self.report, manual_values={("2026-10-13", "412111"): 12.5})
        with self.assertRaises(ValueError):
            render_cell_matrix(self.report, reset_dates={"2026-10-12"})

    def test_local_workbook_has_formulas_preview_and_only_cell_inputs_unlocked(self):
        matrix = render_cell_matrix(self.report, dummy=True)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "dummy.xlsx"
            write_cell_design_workbook(path, [matrix])
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(len(workbook.worksheets), 3)
            formulas, preview = workbook["412110"], workbook["412110 VALUES"]
            self.assertTrue(formulas.protection.sheet)
            self.assertEqual(formulas.freeze_panes, "D2")
            self.assertEqual(formulas["F3"].data_type, "f")
            self.assertEqual(preview["F3"].value, 3000000)
            for merged in formulas.merged_cells.ranges:
                self.assertFalse(merged.min_col <= 3 < merged.max_col)
                self.assertEqual(merged.min_row, 2)
            for column in (7, 14):
                for row in (1, 2, 3, 10):
                    cell = formulas.cell(row, column)
                    self.assertIsNone(cell.value)
                    self.assertEqual(cell.fill.fgColor.rgb, "00FFFFFF")
                    self.assertIsNone(getattr(cell.border.bottom, "style", None))
            self.assertEqual(formulas["F3"].border.right.style, "medium")
            for item in matrix["editable_inputs"]:
                self.assertFalse(formulas.cell(item["row"], item["column"]).protection.locked)
                self.assertTrue(formulas.cell(item["row"], 4).protection.locked)
                self.assertTrue(preview.cell(item["row"], item["column"]).protection.locked)
            workbook.close()
        with self.assertRaises(ValueError):
            write_cell_design_workbook("ignored.xlsx", [render_cell_matrix(self.report)])


class CompactSheetDesignTests(unittest.TestCase):
    def test_capacity_only_grows_and_keeps_a_scroll_buffer(self):
        from finpay_pipeline.cell_sheets import cell_grid_capacity
        self.assertEqual(cell_grid_capacity(99, 14, {"rowCount": 200, "columnCount": 14}),
                         {"rowCount": 200, "columnCount": 100})
        self.assertEqual(cell_grid_capacity(1200, 18, {"rowCount": 1000, "columnCount": 120}),
                         {"rowCount": 1400, "columnCount": 120})
        self.assertEqual(cell_grid_capacity(1200, 18, {"rowCount": 1200, "columnCount": 100}),
                         {"rowCount": 1200, "columnCount": 100})

    def test_shared_details_and_exact_numeric_projection(self):
        report = dummy_cluster_report("412110_B", ("412111", "412112"))
        wide = current_matrix(report, dummy=True, presentation_version=2)
        compact = current_matrix(report, dummy=True)
        self.assertEqual(compact["layout_version"], 3)
        self.assertEqual(len(compact["values"][0]), 14)
        self.assertEqual(compact["values"][1][3:6], ["DEBET", "KREDIT", "SALDO"])
        self.assertEqual(compact["values"][2][5], "=SUM(J3,N3)")
        for old, new in zip(wide["numeric_preview"][2:], compact["numeric_preview"][2:]):
            self.assertEqual(old[:6], new[:6])
            self.assertEqual(old[10:13], new[7:10])
            self.assertEqual(old[17:20], new[11:14])
            self.assertEqual(new.count(new[0]), 1)
            self.assertEqual(new[6], "")
            self.assertEqual(new[10], "")
        self.assertTrue(all(item["column"] in (8, 12) for item in compact["editable_inputs"]))
        self.assertTrue(all(not item["default"].startswith("=K") for item in compact["editable_inputs"]))

    def test_three_cell_layout_has_eighteen_columns_and_reference_palette(self):
        report = dummy_cluster_report("422040_B", ("422041", "422042", "422043"))
        matrix = current_matrix(report, dummy=True)
        self.assertEqual(len(matrix["values"][0]), 18)
        self.assertEqual([b["start_column"] for b in matrix["blocks"]], [4, 8, 12, 16])
        self.assertEqual(matrix["format_palette"]["header"], "1F4E78")
        self.assertEqual(matrix["format_palette"]["cash"], "F6EDFC")
        with TemporaryDirectory() as directory:
            path = Path(directory) / "compact.xlsx"
            write_cell_design_workbook(path, [matrix])
            book = load_workbook(path, data_only=False)
            sheet = book["422040"]
            self.assertEqual(sheet.freeze_panes, "D3")
            self.assertEqual({str(area) for area in sheet.merged_cells.ranges},
                             {"A1:A2", "B1:B2", "C1:C2", "D1:F1", "H1:J1", "L1:N1", "P1:R1"})
            self.assertEqual(sheet["D1"].alignment.horizontal, "center")
            self.assertEqual(sheet["D4"].alignment.horizontal, "right")
            self.assertEqual(sheet["C4"].alignment.horizontal, "left")
            self.assertEqual(sheet["D1"].font.color.rgb, "00FFFFFF")
            self.assertEqual(sheet["D4"].font.color.rgb, "00000000")
            self.assertEqual(sheet["B4"].font.color.rgb, "00000000")
            for section, color in (("MANDIRI:", "00E5EFF9"), ("CASH:", "00F6EDFC")):
                row = next(i + 1 for i, key in enumerate(matrix["row_keys"]) if key.startswith(section))
                self.assertEqual(sheet.cell(row, 4).fill.fgColor.rgb, color)
            book.close()

    def test_reference_section_borders_and_description_padding(self):
        from finpay_pipeline.cell_publication import CellSheetTarget, _format_requests
        report = dummy_cluster_report("412110_B", ("412111", "412112"))
        matrix = current_matrix(report, dummy=True)
        target = CellSheetTarget("dummy", 7, "writer@example.test", ("writer@example.test",))
        requests = _format_requests(target, matrix)
        width = next(item["updateDimensionProperties"]["properties"]["pixelSize"] for item in requests
                     if item.get("updateDimensionProperties", {}).get("range", {}).get("startIndex") == 2)
        self.assertEqual(width, 430)
        self.assertTrue(any(item.get("repeatCell", {}).get("cell", {}).get("userEnteredFormat", {}).get("wrapStrategy") == "CLIP"
                            for item in requests))
        for section in ("DETAIL:", "MANDIRI:", "CASH:"):
            row = next(i for i, key in enumerate(matrix["row_keys"]) if key.startswith(section))
            strong = [item["updateBorders"] for item in requests if item.get("updateBorders", {}).get("top", {}).get("style") == "SOLID_MEDIUM"
                      and item["updateBorders"]["range"]["startRowIndex"] == row]
            self.assertEqual(bool(strong), section != "CASH:")
        for label in ("MANDIRI:net:Total", "CASH:cash:SELISIH"):
            row = matrix["row_keys"].index(label)
            strong = [item["updateBorders"]["range"] for item in requests
                      if item.get("updateBorders", {}).get("top", {}).get("style") == "SOLID_MEDIUM"
                      and item["updateBorders"]["range"]["startRowIndex"] == row]
            self.assertEqual([(g["startColumnIndex"], g["endColumnIndex"]) for g in strong], [(3, 4), (7, 8), (11, 12)])


    def test_v4_removes_only_aggregate_and_preserves_each_cell_numeric_projection(self):
        for account, codes in (("412110_B", ("412111", "412112")),
                               ("422040_B", ("422041", "422042", "422043"))):
            report = dummy_cluster_report(account, codes)
            old = current_matrix(report, dummy=True)
            new = current_matrix(report, dummy=True, presentation_version=4)
            self.assertEqual(new["layout_version"], 4)
            self.assertEqual([block["cell_code"] for block in new["blocks"]], list(codes))
            self.assertEqual([block["start_column"] for block in new["blocks"]], [4, 8, 12][:len(codes)])
            self.assertEqual(len(new["values"][0]), 4 * len(codes) + 2)
            for before, after in zip(old["numeric_preview"], new["numeric_preview"]):
                self.assertEqual(before[:3], after[:3])
                for index in range(len(codes)):
                    self.assertEqual(before[7 + index * 4:10 + index * 4], after[3 + index * 4:6 + index * 4])
            for entry in new["editable_inputs"]:
                self.assertEqual(entry["column"], 4 + 4 * codes.index(entry["cell_code"]))
            with TemporaryDirectory() as directory:
                path = Path(directory) / "cell-only.xlsx"
                write_cell_design_workbook(path, [new])
                book = load_workbook(path, data_only=False)
                self.assertEqual(book.worksheets[0]["B6"].value, len(codes))
                self.assertEqual(book[account[:-2]].freeze_panes, "D3")
                self.assertEqual(book[account[:-2]].max_column, 4 * len(codes) + 2)
                self.assertTrue(all("TOTAL" not in str(cell.value).upper()
                                    for cell in book[account[:-2]][1]))
                book.close()

    def test_dummy_generator_explicit_v4_keeps_full_cell_count(self):
        from finpay_pipeline.cell_design import generate_dummy_cell_design
        with TemporaryDirectory() as directory:
            result = generate_dummy_cell_design(directory, presentation_version=4)
            self.assertEqual((result["clusters"], result["cells"]), (8, 18))

    def test_v4_titles_match_wide_title_columns_for_two_and_three_cells(self):
        for account, codes in (("412110_B", ("412111", "412112")),
                               ("422040_B", ("422041", "422042", "422043"))):
            for dummy in (False, True):
                matrix = current_matrix(dummy_cluster_report(account, codes), dummy=dummy, presentation_version=4)
                expected = [("DUMMY | " if dummy else "") + f"C{index} | CELL {code}"
                            for index, code in enumerate(codes, 1)]
                for kind in ("values", "numeric_preview"):
                    self.assertEqual([matrix[kind][0][block["start_column"] - 1]
                                      for block in matrix["blocks"]], expected)


    def test_v5_bank_inputs_stay_debet_but_reference_next_day_credit(self):
        report = dummy_cluster_report("412110_B", ("412111", "412112"))
        matrix = current_matrix(report, dummy=True, presentation_version=5)
        for entry in matrix["editable_inputs"]:
            index = next(i for i, day in enumerate(report["days"]) if day["report_date"] == entry["report_date"])
            self.assertIn(entry["column"], (4, 8))
            if index == len(report["days"]) - 1:
                self.assertEqual(entry["default"], "PENDING NEXT DAY")
                continue
            next_day = report["days"][index + 1]
            cell = next(c for c in next_day["cells"] if c["cell_code"] == entry["cell_code"])
            next_area = matrix["date_ranges"][index + 1]
            target = next(row for row in range(next_area["start_row"], next_area["end_row"] + 1)
                          if matrix["row_keys"][row - 1] == "DETAIL:CASHOUT APOLLO")
            letter = "E" if entry["column"] == 4 else "I"
            self.assertEqual(entry["default"], f"={letter}{target}")
            self.assertEqual(matrix["numeric_preview"][entry["row"] - 1][entry["column"] - 1],
                             cell["measures"].get("CASHOUT APOLLO", {}).get("kredit", 0))


if __name__ == "__main__":
    unittest.main()
