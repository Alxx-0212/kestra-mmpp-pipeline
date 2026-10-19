from datetime import date
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from threading import Barrier, Lock, get_ident
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from finpay_pipeline.loading import load_and_validate_schema
from finpay_pipeline.digipos_scraper import (
    CLUSTERS,
    HEADERS,
    _download_via_datatables_xhr,
    _download_cluster,
    _datatables_row_to_cells,
    _read_xlsx_rows,
    _replace_datatables_length,
    _rows_from_datatables_response,
    _write_xhr_export,
    expected_source_filename,
    parse_datatable_total,
    parse_report_date,
    scrape_all_clusters,
    validate_export_rows,
)


def sample_rows(transaction_date="28-09-2026 12:30:00", nomor_rs="628123456789"):
    return [
        ["DigiPOS CMS - Riwayat Saldo"],
        list(HEADERS),
        [1, transaction_date, "TX-1", 100, 10, 0, 110, "CREDIT", "RECHARGE", nomor_rs, "OK"],
    ]


class DigiPosExportValidationTest(unittest.TestCase):
    def test_report_date_and_canonical_finpay_filename(self):
        self.assertEqual(parse_report_date("2026-09-28"), date(2026, 9, 28))
        self.assertEqual(
            expected_source_filename("411311", "2026-09-28"),
            "finpay-411311(28-09-2026to28-09-2026).xlsx",
        )

    def test_all_six_expected_clusters_are_configured(self):
        self.assertEqual(CLUSTERS, ("411311", "421306", "421307", "421315", "421318", "421320"))

    def test_accepts_matching_header_count_date_and_identifier(self):
        self.assertEqual(validate_export_rows(sample_rows(), "2026-09-28", 1), 1)

    def test_rejects_row_count_mismatch(self):
        with self.assertRaisesRegex(ValueError, "row mismatch"):
            validate_export_rows(sample_rows(), "2026-09-28", 2)

    def test_rejects_wrong_transaction_date(self):
        with self.assertRaisesRegex(ValueError, "outside requested date"):
            validate_export_rows(sample_rows("27-09-2026 23:59:59"), "2026-09-28", 1)

    def test_rejects_header_only_export(self):
        with self.assertRaisesRegex(ValueError, "header-only"):
            validate_export_rows([list(HEADERS)], "2026-09-28", 0)

    def test_rejects_scientific_notation_nomor_rs(self):
        with self.assertRaisesRegex(ValueError, "scientific notation"):
            validate_export_rows(sample_rows(nomor_rs="6.281234E+11"), "2026-09-28", 1)

    def test_rejects_missing_or_extra_schema_columns(self):
        rows = sample_rows()
        rows[1] = rows[1][:-1]
        rows[2] = rows[2][:-1]
        with self.assertRaisesRegex(ValueError, "Unexpected FinPay XLSX columns"):
            validate_export_rows(rows, "2026-09-28", 1)

    def test_parses_datatables_thousands_separators(self):
        self.assertEqual(parse_datatable_total("Showing 1 to 10 of 3,111 entries"), 3111)

    def test_native_exports_reject_missing_identifiers_and_uningestible_amounts(self):
        for column, value in ((2, ""), (3, None), (4, "invalid"), (5, "NaN"), (6, "")):
            with self.subTest(column=column, value=value):
                rows = sample_rows()
                rows[2][column] = value
                with self.assertRaises(ValueError):
                    validate_export_rows(rows, "2026-09-28", 1)


class DigiPosParallelAcquisitionTest(unittest.TestCase):
    def test_bounded_workers_overlap_and_manifest_preserves_cluster_order(self):
        for workers in (1, 3, 6):
            with self.subTest(workers=workers):
                barrier = Barrier(workers, timeout=10)
                lock = Lock()
                active = 0
                peak = 0

                def download(cluster_id, report_date, password, output_dir, max_attempts):
                    nonlocal active, peak
                    self.assertEqual(report_date, date(2026, 9, 28))
                    self.assertEqual(max_attempts, 5)
                    with lock:
                        active += 1
                        peak = max(peak, active)
                    try:
                        barrier.wait()
                        return {"cluster_id": cluster_id}
                    finally:
                        with lock:
                            active -= 1

                with TemporaryDirectory() as tmp_dir:
                    with patch("finpay_pipeline.digipos_scraper._download_cluster", side_effect=download) as mocked:
                        kwargs = {} if workers == 3 else {"max_workers": workers}
                        manifest = scrape_all_clusters("2026-09-28", "test-only", tmp_dir, **kwargs)
                    self.assertTrue((Path(tmp_dir) / "digipos_manifest.json").exists())

                self.assertEqual(mocked.call_count, len(CLUSTERS))
                self.assertEqual(peak, workers)
                self.assertTrue(manifest["all_clusters_validated"])
                self.assertEqual([item["cluster_id"] for item in manifest["files"]], list(CLUSTERS))

    def test_failed_cluster_blocks_manifest_and_preserves_other_downloads(self):
        def download(cluster_id, report_date, password, output_dir, max_attempts):
            if cluster_id == CLUSTERS[0]:
                raise RuntimeError("cluster retries exhausted")
            _write_xhr_export(output_dir / f"{cluster_id}.xlsx", sample_rows()[2:])
            return {"cluster_id": cluster_id}

        with TemporaryDirectory() as tmp_dir:
            directory = Path(tmp_dir)
            manifest_path = directory / "digipos_manifest.json"
            manifest_path.write_text("stale manifest", encoding="utf-8")
            with patch("finpay_pipeline.digipos_scraper._download_cluster", side_effect=download) as mocked:
                with self.assertRaisesRegex(RuntimeError, "retries exhausted"):
                    scrape_all_clusters("2026-09-28", "test-only", tmp_dir)
            self.assertEqual(mocked.call_count, len(CLUSTERS))
            self.assertFalse(manifest_path.exists())
            self.assertFalse((directory / f"{CLUSTERS[0]}.xlsx").exists())
            for cluster_id in CLUSTERS[1:]:
                self.assertTrue((directory / f"{cluster_id}.xlsx").exists())

    def test_rejects_invalid_worker_counts_before_touching_artifacts(self):
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "digipos_manifest.json"
            path.write_text("existing manifest", encoding="utf-8")
            for workers in (0, 7, -1, 1.5, "3", True):
                with self.subTest(workers=workers):
                    with self.assertRaisesRegex(ValueError, "max_workers must be an integer"):
                        scrape_all_clusters("2026-09-28", "test-only", tmp_dir, max_workers=workers)
                    self.assertEqual(path.read_text(encoding="utf-8"), "existing manifest")

    def test_playwright_and_browser_lifecycle_stays_on_worker_threads(self):
        barrier = Barrier(3, timeout=10)
        browsers = []
        lock = Lock()

        class Browser:
            def __init__(self):
                self.owner = get_ident()
                self.closed_on = None
                with lock:
                    browsers.append(self)

            def close(self):
                self.closed_on = get_ident()

        def launch(headless):
            return Browser()

        def download(browser, cluster_id, report_date, password, output_dir, max_attempts):
            self.assertEqual(browser.owner, get_ident())
            barrier.wait()
            return {"cluster_id": cluster_id}

        def playwright_context():
            context = MagicMock()
            context.__enter__.return_value = SimpleNamespace(chromium=SimpleNamespace(launch=launch))
            return context

        sync_api = ModuleType("playwright.sync_api")
        sync_api.sync_playwright = playwright_context
        with TemporaryDirectory() as tmp_dir:
            with patch.dict(sys.modules, {"playwright.sync_api": sync_api}):
                with patch("finpay_pipeline.digipos_scraper._download_one_cluster", side_effect=download):
                    scrape_all_clusters("2026-09-28", "test-only", tmp_dir)

        self.assertEqual(len(browsers), len(CLUSTERS))
        self.assertEqual(len({browser.owner for browser in browsers}), 3)
        for browser in browsers:
            self.assertNotEqual(browser.owner, get_ident())
            self.assertEqual(browser.closed_on, browser.owner)

    def test_worker_closes_browser_when_cluster_fails(self):
        browser = MagicMock()
        context = MagicMock()
        context.__enter__.return_value.chromium.launch.return_value = browser
        sync_api = ModuleType("playwright.sync_api")
        sync_api.sync_playwright = lambda: context
        with patch.dict(sys.modules, {"playwright.sync_api": sync_api}):
            with patch("finpay_pipeline.digipos_scraper._download_one_cluster", side_effect=RuntimeError("failed")):
                with self.assertRaisesRegex(RuntimeError, "failed"):
                    _download_cluster(CLUSTERS[0], date(2026, 9, 28), "test-only", Path("."), 5)
        browser.close.assert_called_once_with()
        context.__exit__.assert_called_once()


class DigiPosXhrFallbackTest(unittest.TestCase):
    @staticmethod
    def source_row(**overrides):
        return {
            "tanggal": "28-09-2026 12:30:00",
            "wstransferid": "000123",
            "init_amount": "1000",
            "amount_in": "50",
            "amount_out": "0",
            "last_amount": "1050",
            "transaction_type": "Credit",
            "wstransfertype": "RECHARGE",
            "sumber1": "490812345678",
            "wstransferdesc": "source remarks",
            **overrides,
        }

    def test_exact_identifiers_and_phone_survive_xlsx_and_daily_loading(self):
        cases = (
            ["000123", "123", "000123FEE", "12345678901234567890"],
            ["000123", "123", "000456"],
        )
        for identifiers in cases:
            with self.subTest(identifiers=identifiers), TemporaryDirectory() as tmp_dir:
                path = Path(tmp_dir) / "fallback.xlsx"
                rows = [
                    _datatables_row_to_cells(self.source_row(wstransferid=value), index)
                    for index, value in enumerate(identifiers, 1)
                ]
                _write_xhr_export(path, rows)
                loaded = load_and_validate_schema(str(path))
                self.assertEqual(loaded["Transaction ID"].tolist(), identifiers)
                self.assertEqual(loaded["Nomor RS"].tolist(), ["0812345678"] * len(identifiers))

    def test_missing_identifiers_are_rejected_instead_of_fabricated(self):
        for value in (None, "", "   "):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "Transaction ID"):
                    _datatables_row_to_cells(self.source_row(wstransferid=value), 1)

    def test_missing_financial_values_fail_closed_but_explicit_zero_is_valid(self):
        for field in ("init_amount", "amount_in", "amount_out", "last_amount"):
            for value in (None, "", "   "):
                with self.subTest(field=field, value=value):
                    with self.assertRaises(ValueError):
                        _datatables_row_to_cells(self.source_row(**{field: value}), 1)
            row = self.source_row()
            del row[field]
            with self.subTest(field=field, absent=True), self.assertRaises(ValueError):
                _datatables_row_to_cells(row, 1)
        self.assertEqual(
            _datatables_row_to_cells(self.source_row(init_amount=0, amount_in="0", amount_out=0, last_amount="0"), 1)[3:7],
            [0, 0, 0, 0],
        )

    def test_equals_prefixed_source_text_is_literal_not_an_excel_formula(self):
        from openpyxl import load_workbook

        for column in (2, 7, 8, 9, 10):
            with self.subTest(column=column), TemporaryDirectory() as tmp_dir:
                path = Path(tmp_dir) / "fallback.xlsx"
                row = _datatables_row_to_cells(self.source_row(), 1)
                row[column] = "=1+1"
                _write_xhr_export(path, [row])
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    self.assertEqual(workbook.active.cell(3, column + 1).data_type, "s")
                finally:
                    workbook.close()
                self.assertEqual(_read_xlsx_rows(path)[2][column], "=1+1")

    def test_numeric_and_date_like_source_text_is_not_coerced(self):
        for file_type in ("xlsx", "csv"):
            with self.subTest(file_type=file_type), TemporaryDirectory() as tmp_dir:
                path = Path(tmp_dir) / f"source.{file_type}"
                row = sample_rows()[2]
                row[2] = "000123"
                row[9] = "0812345678"
                row[10] = "20260101"
                if file_type == "xlsx":
                    _write_xhr_export(path, [row])
                else:
                    import csv
                    with path.open("w", encoding="utf-8", newline="") as handle:
                        writer = csv.writer(handle)
                        writer.writerow(HEADERS)
                        writer.writerow(row)
                loaded = load_and_validate_schema(str(path))
                self.assertEqual(loaded.loc[0, "Transaction ID"], "000123")
                self.assertEqual(loaded.loc[0, "Nomor RS"], "0812345678")
                self.assertEqual(loaded.loc[0, "Remarks"], "20260101")

    def test_rewrites_datatables_page_length_without_changing_other_fields(self):
        post_data = "draw=4&start=0&length=10&search%5Bvalue%5D="
        self.assertEqual(
            _replace_datatables_length(post_data, 123),
            "draw=4&start=0&length=123&search%5Bvalue%5D=",
        )

    def test_requires_datatables_length_field(self):
        with self.assertRaisesRegex(ValueError, "no length field"):
            _replace_datatables_length("draw=4&start=0", 123)

    def test_maps_credit_and_debit_rows_to_export_columns(self):
        row = {
            "tanggal": "28-09-2026 12:30:00",
            "wstransferid": "12345",
            "init_amount": "1,000",
            "amount_in": "50",
            "amount_out": "0",
            "last_amount": "1,050",
            "transaction_type": "Credit",
            "wstransfertype": "RECHARGE",
            "sumber1": "490812345678",
            "wstransferdesc": "source remarks",
        }
        self.assertEqual(
            _datatables_row_to_cells(row, 1),
            [
                1,
                "28-09-2026 12:30:00",
                "12345",
                1000,
                50,
                0,
                1050,
                "Credit",
                "RECHARGE",
                "0812345678",
                "source remarks",
            ],
        )
        row["transaction_type"] = "Debit"
        self.assertEqual(_datatables_row_to_cells(row, 2)[9], "-")

    def test_requires_footer_total_and_returned_rows_to_match(self):
        row = {
            "tanggal": "28-09-2026 12:30:00",
            "wstransferid": "12345",
            "transaction_type": "Credit",
        }
        with self.assertRaisesRegex(ValueError, "footer mismatch"):
            _rows_from_datatables_response(
                {"recordsFiltered": 2, "data": [row]},
                1,
            )
        with self.assertRaisesRegex(ValueError, "row mismatch"):
            _rows_from_datatables_response(
                {"recordsFiltered": 1, "data": []},
                1,
            )

    def test_fallback_xlsx_passes_the_existing_export_validation(self):
        source_row = {
            "tanggal": "28-09-2026 12:30:00",
            "wstransferid": "12345",
            "init_amount": "1000",
            "amount_in": "50",
            "amount_out": "0",
            "last_amount": "1050",
            "transaction_type": "Credit",
            "wstransfertype": "RECHARGE",
            "sumber1": "490812345678",
            "wstransferdesc": "source remarks",
        }
        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "fallback.xlsx"
            _write_xhr_export(path, [_datatables_row_to_cells(source_row, 1)])
            self.assertEqual(
                validate_export_rows(_read_xlsx_rows(path), "2026-09-28", 1),
                1,
            )

    def test_download_rewrites_request_and_writes_response_xlsx(self):
        row = {
            "tanggal": "28-09-2026 12:30:00",
            "wstransferid": "12345",
            "init_amount": "1000",
            "amount_in": "50",
            "amount_out": "0",
            "last_amount": "1050",
            "transaction_type": "Credit",
            "wstransfertype": "RECHARGE",
            "sumber1": "490812345678",
            "wstransferdesc": "source remarks",
        }

        class FakeRoute:
            def __init__(self, request):
                self.request = request
                self.continued_post_data = None
                self.continued_headers = None

            def continue_(self, post_data=None, headers=None):
                self.continued_post_data = post_data
                self.continued_headers = headers

        class FakePage:
            def route(self, pattern, handler):
                self.route_handler = handler

            def unroute(self, pattern):
                self.unrouted_pattern = pattern

            def expect_response(self, predicate, timeout):
                self.response_predicate = predicate
                page = self

                class ResponseContext:
                    def __enter__(self):
                        return self

                    def __exit__(self, *args):
                        return False

                    @property
                    def value(self):
                        return page.response

                return ResponseContext()

            def evaluate(self, script, page_length):
                request = SimpleNamespace(
                    url="https://digipos-cms.finpay.id/deposit/monitoring-riwayat-data",
                    method="POST",
                    post_data="draw=4&start=0&length=10",
                    headers={
                        "content-type": "application/x-www-form-urlencoded",
                        "content-length": "28",
                    },
                )
                self.route = FakeRoute(request)
                self.route_handler(self.route)
                stale_response = SimpleNamespace(
                    url=request.url,
                    request=SimpleNamespace(url=request.url, method="POST"),
                )
                if self.response_predicate(stale_response):
                    raise AssertionError("fallback accepted an earlier unmodified request")
                self.response = SimpleNamespace(
                    url=request.url,
                    request=request,
                    status=200,
                    json=lambda: {"recordsFiltered": 1, "data": [row]},
                )
                if not self.response_predicate(self.response):
                    raise AssertionError("fallback response did not match")

        sync_api = ModuleType("playwright.sync_api")
        sync_api.TimeoutError = type("TimeoutError", (Exception,), {})
        playwright = ModuleType("playwright")
        playwright.sync_api = sync_api
        page = FakePage()

        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "fallback.xlsx"
            with patch.dict(
                sys.modules,
                {"playwright": playwright, "playwright.sync_api": sync_api},
            ):
                row_count = _download_via_datatables_xhr(
                    page,
                    date(2026, 9, 28),
                    1,
                    path,
                )

            self.assertEqual(row_count, 1)
            self.assertEqual(page.route.continued_post_data, "draw=4&start=0&length=1")
            self.assertEqual(page.route.continued_headers["content-length"], "23")
            self.assertEqual(validate_export_rows(_read_xlsx_rows(path), "2026-09-28", 1), 1)


if __name__ == "__main__":
    unittest.main()
