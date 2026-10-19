"""Synthetic bundle-to-report and new workflow public-boundary regressions."""
import ast
import contextlib
import hashlib
import io
import json
import tempfile
import textwrap
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch
from types import SimpleNamespace

import finpay_pipeline
from finpay_pipeline.cell_contracts import OBSERVED_CELL_ROSTER, configured_cells, fingerprint, roster_fingerprint
from finpay_pipeline.cell_calculation import CellReportConfig, calculate_cell_day
from finpay_pipeline.cell_workflow import _bounded_bytes, acquire_cell_bundle, read_cell_configuration, run_cell_reporting
from finpay_pipeline.digipos_scraper import _write_xhr_export
from finpay_pipeline.tests.test_cell_publication import FakeStore


def synthetic_bundle(root, roster=None, *, malformed_last=False, dummy=True):
    roster = roster or OBSERVED_CELL_ROSTER
    directory = Path(root)
    cells = configured_cells(roster)
    configs = [{"account": account, "cell_codes": codes, "first_report_date": "2026-10-13",
                "opening_boundary": "2026-10-13T00:00:00+08:00",
                "opening_balances": {code: 1000 for code in codes}, "roster_version": "DUMMY-ONLY"}
               for account, codes in sorted(roster.items())]
    runtime = {"schema_version": 1, "dummy_data_only": dummy, "accounts": configs}
    config_path = directory / "configuration.json"
    config_path.write_text(json.dumps(runtime), encoding="utf-8")
    entries, contents = [], {}
    for index, cell in enumerate(cells):
        name = cell.account + "-" + cell.cell_code + ".xlsx"
        path = directory / name
        _write_xhr_export(path, [[1, "13-10-2026 12:00:00", "00012345678901234567",
                                  1000, 100, 0, 1100, "Credit", "RECHARGE", "-", "Transaksi Sellthru"]])
        if malformed_last and index == len(cells) - 1:
            from openpyxl import load_workbook
            workbook = load_workbook(path)
            workbook.active.cell(row=3, column=5).value = "not-an-amount"
            workbook.save(path)
            workbook.close()
        content = path.read_bytes()
        contents[name] = content
        entries.append({**cell.to_dict(), "report_date": "2026-10-13", "status": "validated",
                        "validated_rows": 1, "footer_total": 1, "artifact": name,
                        "source_sha256": hashlib.sha256(content).hexdigest()})
    manifest = {"schema_version": 1, "report_date": "2026-10-13", "roster_fingerprint": roster_fingerprint(cells),
                "cells": entries}
    bundle = directory / "cell_sources.zip"
    with zipfile.ZipFile(bundle, "w") as archive:
        archive.writestr("manifest.json", json.dumps(manifest))
        for name, content in contents.items():
            archive.writestr(name, content)
    return bundle, config_path, runtime


class CellWorkflowTests(unittest.TestCase):
    def invoke(self, bundle, config, **kwargs):
        return run_cell_reporting(str(bundle), str(config), hashlib.sha256(config.read_bytes()).hexdigest(),
                                  hashlib.sha256(bundle.read_bytes()).hexdigest(),
                                  output_dir=config.parent / "outputs", **kwargs)

    def test_all_eight_accounts_eighteen_cells_dry_run_has_no_live_factories(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, _ = synthetic_bundle(directory)
            database, google = Mock(), Mock()
            result = self.invoke(bundle, config, session_factory=database, sheet_client_factory=google)
            self.assertEqual(len(result["accounts"]), 8)
            self.assertEqual(sum(item["cells"] for item in result["accounts"]), 18)
            self.assertTrue(all(item["status"] == "DRY_RUN" for item in result["accounts"]))
            database.assert_not_called()
            google.assert_not_called()
            saved = json.loads((config.parent / "outputs/cell_prepared.json").read_text())
            sample = saved["batches"]["412110_B"][0]["raw_rows"][0]
            self.assertEqual(sample["Transaction ID"], "00012345678901234567")
            self.assertNotIn("No", sample)

    def test_bad_last_account_blocks_every_write_not_only_that_account(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, _ = synthetic_bundle(directory, malformed_last=True, dummy=False)
            database, google = Mock(), Mock()
            with self.assertRaises(ValueError):
                self.invoke(bundle, config, dry_run=False, session_factory=database, sheet_client_factory=google,
                            approved_configuration_sha256=hashlib.sha256(config.read_bytes()).hexdigest())
            database.assert_not_called()
            google.assert_not_called()

    def test_dummy_configuration_cannot_publish(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, _ = synthetic_bundle(directory, {"412110_B": ("412111", "412112")})
            database = Mock()
            with self.assertRaisesRegex(ValueError, "dummy configurations"):
                self.invoke(bundle, config, dry_run=False, session_factory=database,
                            approved_configuration_sha256=hashlib.sha256(config.read_bytes()).hexdigest())
            database.assert_not_called()

    def test_database_only_processes_accounts_serially_without_google(self):
        roster = {"412110_B": ("412111", "412112"), "422040_B": ("422041", "422042", "422043")}
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, runtime = synthetic_bundle(directory, roster, dummy=False)
            target = {"host": "dummy.test", "port": 5432, "database": "dummy_finpay",
                      "schema": "dummy_schema", "user": "dummy_role"}
            runtime.update(database_target=target, approved_database_targets=[target])
            config.write_text(json.dumps(runtime), encoding="utf-8")
            order = []

            @contextlib.contextmanager
            def session(target, approved, cluster, sheet, **kwargs):
                self.assertEqual(sheet, "")
                self.assertFalse(kwargs["write_gsheet"])
                order.append("enter:" + cluster)
                store = FakeStore()
                store.cluster_id = cluster
                store.target_id = "database-only:" + cluster
                yield store
                order.append("exit:" + cluster)

            google = Mock()
            result = self.invoke(bundle, config, dry_run=False, write_gsheet=False,
                                 database_password="synthetic", session_factory=session, sheet_client_factory=google,
                                 approved_configuration_sha256=hashlib.sha256(config.read_bytes()).hexdigest())
            self.assertEqual(order, ["enter:412110", "exit:412110", "enter:422040", "exit:422040"] * 2)
            self.assertEqual(len(result["accounts"]), 2)
            google.assert_not_called()

    def test_later_account_missing_history_blocks_all_writes_in_global_preflight(self):
        roster = {"412110_B": ("412111",), "422040_B": ("422041", "422042")}
        with tempfile.TemporaryDirectory() as directory:
            bundle, config_path, runtime = synthetic_bundle(directory, roster, dummy=False)
            target = {"host": "dummy.test", "port": 5432, "database": "dummy_finpay",
                      "schema": "dummy_schema", "user": "dummy_role"}
            runtime.update(database_target=target, approved_database_targets=[target])
            config_path.write_text(json.dumps(runtime))
            stores = []

            @contextlib.contextmanager
            def session(target, allowed, cluster, sheet, **kwargs):
                store = FakeStore()
                store.cluster_id = cluster
                stores.append(store)
                if cluster == "422040":
                    config = CellReportConfig(**{**runtime["accounts"][1], "cell_codes": tuple(roster["422040_B"])})
                    cell = config.cells[0]
                    row = {"No": 1, "Transaction Date": "2026-10-14 12:00:00", "Transaction ID": "DUMMY",
                           "Saldo Awal": 1000, "Kredit": 100, "Debet": 0, "Saldo Akhir": 1100,
                           "Transaction Type": "Credit", "Transaction": "RECHARGE", "Nomor RS": "-",
                           "Remarks": "Transaksi Sellthru"}
                    entry = {**cell.to_dict(), "report_date": "2026-10-14", "status": "validated",
                             "artifact": "dummy.xlsx", "source_sha256": fingerprint([row]),
                             "validated_rows": 1, "footer_total": 1}
                    store.facts[("2026-10-14", cell.cell_code)] = calculate_cell_day(cell, entry, [row], config)
                yield store

            google = Mock()
            with self.assertRaisesRegex(ValueError, "missing cell coverage"):
                self.invoke(bundle, config_path, dry_run=False, write_gsheet=False, database_password="synthetic",
                            session_factory=session, sheet_client_factory=google,
                            approved_configuration_sha256=hashlib.sha256(config_path.read_bytes()).hexdigest())
            self.assertEqual(len(stores), 2)
            self.assertTrue(all("commit" not in store.events for store in stores))
            google.assert_not_called()

    def test_wrong_config_hash_rejected_before_any_client(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, _ = synthetic_bundle(directory, {"412110_B": ("412111",)})
            with self.assertRaisesRegex(ValueError, "checksum"):
                read_cell_configuration(config, "0" * 64)

    def test_self_approved_target_and_matching_input_hash_cannot_authorize_a_connection(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, runtime = synthetic_bundle(directory, {"412110_B": ("412111",)}, dummy=False)
            target = {"host": "attacker.invalid", "port": 5432, "database": "dummy",
                      "schema": "public", "user": "dummy_role"}
            runtime.update(database_target=target, approved_database_targets=[target])
            config.write_text(json.dumps(runtime))
            for trusted in (None, "0" * 64):
                with self.subTest(trusted=trusted):
                    database, google = Mock(), Mock()
                    with self.assertRaisesRegex(ValueError, "deployment-approved"):
                        self.invoke(bundle, config, dry_run=False, write_gsheet=False,
                                    database_password="synthetic-only", session_factory=database,
                                    sheet_client_factory=google, approved_configuration_sha256=trusted)
                    database.assert_not_called()
                    google.assert_not_called()

    def test_acquisition_also_requires_deployment_pin_before_source_calls(self):
        with tempfile.TemporaryDirectory() as directory:
            _, config, _ = synthetic_bundle(directory, {"412110_B": ("412111",)})
            with patch("finpay_pipeline.cell_workflow.scrape_cell_roster") as scrape:
                with self.assertRaisesRegex(ValueError, "deployment-approved"):
                    acquire_cell_bundle("2026-10-13", config, hashlib.sha256(config.read_bytes()).hexdigest(),
                                        directory, password="synthetic")
                scrape.assert_not_called()

    def test_size_rejection_precedes_read_and_handles_file_growth(self):
        with patch.object(Path, "stat", return_value=SimpleNamespace(st_size=100)), \
                patch.object(Path, "open") as opened:
            with self.assertRaisesRegex(ValueError, "size limit"):
                _bounded_bytes("dummy.json", 10)
            opened.assert_not_called()
        with patch.object(Path, "stat", return_value=SimpleNamespace(st_size=1)), \
                patch.object(Path, "open", return_value=io.BytesIO(b"x" * 11)):
            with self.assertRaisesRegex(ValueError, "grew"):
                _bounded_bytes("dummy.json", 10)

    def test_acquisition_rejects_oversized_artifact_or_bundle_before_packaging(self):
        for artifact_limit, bundle_limit in ((1, 256 * 1024 * 1024), (32 * 1024 * 1024, 100)):
            with self.subTest(artifact=artifact_limit), tempfile.TemporaryDirectory() as directory:
                bundle, config, _ = synthetic_bundle(directory, {"412110_B": ("412111",)}, dummy=False)
                with zipfile.ZipFile(bundle) as archive:
                    manifest = json.loads(archive.read("manifest.json"))
                dated = Path(directory) / "2026-10-13"
                dated.mkdir()
                for entry in manifest["cells"]:
                    (Path(directory) / entry["artifact"]).replace(dated / entry["artifact"])
                with patch("finpay_pipeline.cell_workflow.scrape_cell_roster", return_value=manifest), \
                        patch("finpay_pipeline.cell_workflow.MAX_CELL_ARTIFACT_BYTES", artifact_limit), \
                        patch("finpay_pipeline.cell_workflow.MAX_BUNDLE_BYTES", bundle_limit), \
                        patch("finpay_pipeline.cell_workflow.sha256_file_bounded") as hashed:
                    with self.assertRaisesRegex(ValueError, "acquisition size limit"):
                        acquire_cell_bundle("2026-10-13", config, hashlib.sha256(config.read_bytes()).hexdigest(),
                                            directory, password="synthetic", approved_configuration_sha256=hashlib.sha256(config.read_bytes()).hexdigest())
                    hashed.assert_not_called()
    def test_archive_paths_and_unexpected_files_are_rejected_without_extraction(self):
        for name in ("../escape.xlsx", "unexpected.xlsx"):
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                bundle, config, _ = synthetic_bundle(directory, {"412110_B": ("412111",)})
                with zipfile.ZipFile(bundle, "a") as archive:
                    archive.writestr(name, "DUMMY")
                with self.assertRaises(ValueError):
                    self.invoke(bundle, config)

    def test_each_artifact_hash_is_checked_beyond_bundle_hash(self):
        with tempfile.TemporaryDirectory() as directory:
            bundle, config, _ = synthetic_bundle(directory, {"412110_B": ("412111",)})
            with zipfile.ZipFile(bundle) as archive:
                contents = {name: archive.read(name) for name in archive.namelist()}
            manifest = json.loads(contents["manifest.json"])
            manifest["cells"][0]["source_sha256"] = "0" * 64
            contents["manifest.json"] = json.dumps(manifest).encode()
            with zipfile.ZipFile(bundle, "w") as archive:
                for name, content in contents.items():
                    archive.writestr(name, content)
            with self.assertRaisesRegex(ValueError, "XLSX checksum"):
                self.invoke(bundle, config)

    def test_new_flows_keep_public_imports_scripts_and_no_notifications(self):
        directory = Path(__file__).resolve().parents[1]
        for filename in ("finpay_digipos_cell_acquisition_v2.yml", "finpay_cell_reporting_v1.yml"):
            text = (directory / "workflows" / filename).read_text()
            self.assertNotIn("TELEGRAM", text)
            self.assertNotIn("finpay_daily_pipeline_v5", text)
            self.assertIn("behavior: QUEUE", text)
            lines = text.splitlines()
            start = next(index for index, line in enumerate(lines) if line.strip() == "script: |")
            indent = len(lines[start]) - len(lines[start].lstrip())
            script = []
            for line in lines[start + 1:]:
                if line.strip() and len(line) - len(line.lstrip()) <= indent:
                    break
                script.append(line)
            parsed = ast.parse(textwrap.dedent("\n".join(script)))
            for node in ast.walk(parsed):
                if isinstance(node, ast.ImportFrom) and node.module == "finpay_pipeline":
                    for imported in node.names:
                        self.assertTrue(hasattr(finpay_pipeline, imported.name), imported.name)
            self.assertIn("except Exception as exc", text)
        acquisition = (directory / "workflows/finpay_digipos_cell_acquisition_v2.yml").read_text()
        reporting = (directory / "workflows/finpay_cell_reporting_v1.yml").read_text()
        for text in (acquisition, reporting):
            self.assertIn("secret('FINPAY_CELL_CONFIG_SHA256')", text)
            self.assertIn('approved_configuration_sha256=os.environ["CELL_APPROVED_CONFIGURATION_SHA256"]', text)
        self.assertIn('cron: "0 5 * * *"', acquisition)
        self.assertIn("timezone: Asia/Makassar", acquisition)
        self.assertIn("disabled: true", acquisition)
        self.assertIn("dateAdd(-1, 'DAYS')", acquisition)

    def test_runtime_copy_excludes_local_generated_data_and_configuration(self):
        package = Path(__file__).resolve().parents[1]
        for name in ("Dockerfile", "Dockerfile.cell-scraper"):
            text = (package / name).read_text()
            self.assertNotIn("COPY finpay_pipeline ./finpay_pipeline", text)
            self.assertIn("COPY finpay_pipeline/*.py ./finpay_pipeline/", text)
            self.assertIn("COPY finpay_pipeline/migrations ./finpay_pipeline/migrations", text)
            self.assertIn("COPY finpay_pipeline/queries ./finpay_pipeline/queries", text)


if __name__ == "__main__":
    unittest.main()
