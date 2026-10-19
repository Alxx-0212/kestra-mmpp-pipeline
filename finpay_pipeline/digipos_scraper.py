"""Kestra-side DigiPOS daily XLSX acquisition and validation."""
from __future__ import annotations

import hashlib
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from pathlib import Path


BASE_URL = "https://digipos-cms.finpay.id"
ENDPOINT_PATH = "/deposit/monitoring-riwayat-data"
CLUSTERS = ("411311", "421306", "421307", "421315", "421318", "421320")
HEADERS = (
    "No",
    "Transaction Date",
    "Transaction ID",
    "Saldo Awal",
    "Kredit",
    "Debet",
    "Saldo Akhir",
    "Transaction Type",
    "Transaction",
    "Nomor RS",
    "Remarks",
)
MAX_TABLE_ROWS = 20_000
XLSX_TITLE = "CMS DigiPOS - Riwayat Saldo"
XLSX_WIDTHS = {
    "A": 6.0, "B": 25.65, "C": 37.8, "D": 14.85, "E": 13.5, "F": 14.85,
    "G": 14.85, "H": 21.6, "I": 18.9, "J": 27.0, "K": 54.0,
}


def parse_report_date(value: str) -> date:
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError("report_date must be YYYY-MM-DD") from exc


def expected_source_filename(cluster_id: str, report_date: str) -> str:
    day = parse_report_date(report_date).strftime("%d-%m-%Y")
    return f"finpay-{cluster_id}({day}to{day}).xlsx"


def normalize_export_rows(rows):
    for index, row in enumerate(rows):
        cells = ["" if value is None else str(value).strip() for value in row]
        if len(cells) >= 2 and cells[0] == "No" and cells[1] == "Transaction Date":
            return rows[index:]
    raise ValueError("Excel export does not contain the expected FinPay header")


def _parse_transaction_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in (
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid Transaction Date in Excel export: {text!r}")


def parse_datatable_total(info_text: str) -> int:
    match = re.search(r"of\s+([\d,]+)\s+entries", info_text or "")
    if not match:
        raise ValueError(f"Could not parse DigiPOS DataTables total: {info_text!r}")
    return int(match.group(1).replace(",", ""))


def _replace_datatables_length(post_data: str, length: int) -> str:
    updated, replacements = re.subn(
        r"(^|&)length=-?\d+(?=&|$)",
        lambda match: f"{match.group(1)}length={length}",
        post_data,
        count=1,
    )
    if replacements != 1:
        raise ValueError("DigiPOS DataTables request has no length field")
    return updated


def _export_integer(value, field_name: str) -> int:
    text = str(value if value is not None else "").strip().replace(",", "")
    if not re.fullmatch(r"-?\d+", text):
        raise ValueError(f"DigiPOS XHR returned a non-integer {field_name}")
    return int(text)


def _datatables_row_to_cells(row: dict, index: int) -> list:
    transaction_id = str(row.get("wstransferid") or "").strip()
    if not transaction_id:
        raise ValueError("DigiPOS XHR returned a missing Transaction ID")

    transaction_type = str(row.get("transaction_type") or "").strip()
    if transaction_type == "Credit":
        nomor_rs = str(row.get("sumber1") or "").strip()
        if nomor_rs.startswith("49"):
            nomor_rs = nomor_rs[2:]
        nomor_rs = nomor_rs or "-"
    else:
        nomor_rs = "-"

    return [
        index,
        str(row.get("tanggal") or "").strip(),
        transaction_id,
        _export_integer(row.get("init_amount"), "Saldo Awal"),
        _export_integer(row.get("amount_in"), "Kredit"),
        _export_integer(row.get("amount_out"), "Debet"),
        _export_integer(row.get("last_amount"), "Saldo Akhir"),
        transaction_type,
        str(row.get("wstransfertype") or "").strip(),
        nomor_rs,
        str(row.get("wstransferdesc") or "").strip(),
    ]


def _rows_from_datatables_response(payload: dict, expected_total: int) -> list[list]:
    if not isinstance(payload, dict):
        raise ValueError("DigiPOS XHR response is not an object")
    try:
        filtered_total = int(payload.get("recordsFiltered"))
    except (TypeError, ValueError) as exc:
        raise ValueError("DigiPOS XHR response has no valid recordsFiltered total") from exc

    data = payload.get("data")
    if filtered_total != expected_total:
        raise ValueError(
            f"DigiPOS XHR footer mismatch: DataTables={expected_total}, "
            f"recordsFiltered={filtered_total}"
        )
    if not isinstance(data, list) or len(data) != expected_total:
        actual_total = len(data) if isinstance(data, list) else "missing"
        raise ValueError(
            f"DigiPOS XHR row mismatch: DataTables={expected_total}, rows={actual_total}"
        )
    return [_datatables_row_to_cells(row, index) for index, row in enumerate(data, 1)]


def _write_xhr_export(path: Path, rows: list[list]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.cell(row=1, column=1, value=XLSX_TITLE)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    worksheet.append(list(HEADERS))
    for row in rows:
        worksheet.append(row)
        for column in (2, 3, 8, 9, 10, 11):
            # Preserve exact identifiers and literal source text, never formulas.
            cell = worksheet.cell(row=worksheet.max_row, column=column)
            cell.data_type = "s"

    for row_number in range(3, worksheet.max_row + 1):
        for column in (4, 5, 6, 7):
            value = worksheet.cell(row=row_number, column=column).value
            worksheet.cell(row=row_number, column=column).number_format = (
                "#,##0" if abs(value) >= 1000 else "0"
            )
        worksheet.cell(row=row_number, column=3).number_format = "@"
    for column, width in XLSX_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    workbook.save(path)
    workbook.close()


def validate_export_rows(rows, report_date: str, expected_total: int) -> int:
    """Reject incomplete or uningestible exports before any reporting child starts."""
    target = parse_report_date(report_date)
    normalized = normalize_export_rows(rows)
    header = ["" if value is None else str(value).strip() for value in normalized[0]]
    if len(header) != len(HEADERS) or set(header) != set(HEADERS):
        raise ValueError(f"Unexpected FinPay XLSX columns: {header!r}")

    date_index = header.index("Transaction Date")
    id_index = header.index("Transaction ID")
    rs_index = header.index("Nomor RS")
    data_rows = [
        row for row in normalized[1:]
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not data_rows:
        raise ValueError("DigiPOS returned an empty/header-only daily export")
    if len(data_rows) != expected_total:
        raise ValueError(
            f"XLSX row mismatch: DataTables={expected_total}, export={len(data_rows)}"
        )
    if expected_total > MAX_TABLE_ROWS:
        raise ValueError(
            f"DigiPOS returned {expected_total} rows, beyond the known per-query "
            f"limit of {MAX_TABLE_ROWS}; split the report date window before processing"
        )

    scientific = []
    for row_number, row in enumerate(data_rows, start=1):
        if len(row) < len(HEADERS):
            raise ValueError(f"XLSX row {row_number} has fewer than {len(HEADERS)} columns")
        if _parse_transaction_date(row[date_index]) != target:
            raise ValueError(f"XLSX row {row_number} is outside requested date {target}")
        if not str(row[id_index] if row[id_index] is not None else "").strip():
            raise ValueError(f"XLSX row {row_number} has a missing Transaction ID")
        for name in ("No", "Saldo Awal", "Kredit", "Debet", "Saldo Akhir"):
            _export_integer(row[header.index(name)], name)
        rs_number = str(row[rs_index] or "").strip()
        if rs_number and rs_number != "-" and re.search(r"\d+\.?\d*E\+?\d+", rs_number, re.I):
            scientific.append(rs_number)
    if scientific:
        raise ValueError(f"Nomor RS values were serialized in scientific notation: {scientific[:5]}")
    return len(data_rows)


def _read_xlsx_rows(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return [list(row) for row in workbook.active.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _set_search_filters(page, report_date: date) -> None:
    display_date = report_date.strftime("%d %B %Y")
    api_date = report_date.isoformat()
    page.evaluate(
        """([startVal, endVal, apiStart, apiEnd]) => {
            const setValue = (selector, value) => {
                const element = document.querySelector(selector);
                if (!element) throw new Error(`Missing DigiPOS date input: ${selector}`);
                const setter = Object.getOwnPropertyDescriptor(
                    HTMLInputElement.prototype, 'value'
                ).set;
                setter.call(element, value);
                element.dispatchEvent(new Event('input', {bubbles: true}));
                element.dispatchEvent(new Event('change', {bubbles: true}));
            };
            setValue('input[name="date1"]', startVal);
            setValue('input[name="date2"]', endVal);
            setValue('#dateStart', apiStart);
            setValue('#dateEnd', apiEnd);
        }""",
        [
            f"{display_date} 00:00:00",
            f"{display_date} 23:59:59",
            f"{api_date} 00:00:00",
            f"{api_date} 23:59:59",
        ],
    )


def _table_total_for_date(page, report_date: date) -> int:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

    server_busy = []

    def watch_response(response):
        if "monitoring-riwayat-data" not in response.url:
            return
        try:
            payload = response.json()
        except Exception:
            return
        if isinstance(payload, dict) and str(payload.get("statusCode")) == "05":
            server_busy.append(str(payload.get("statusDesc", "DigiPOS server busy")))

    page.on("response", watch_response)
    try:
        page.goto(f"{BASE_URL}/deposit/monitoring-riwayat", wait_until="networkidle", timeout=45000)
        page.wait_for_selector("input[name=date1]", timeout=20000)
        page.wait_for_function(
            "() => typeof $ !== 'undefined' && $.fn.DataTable.isDataTable('#dataInput')",
            timeout=20000,
        )
        for search_attempt in range(1, 4):
            _set_search_filters(page, report_date)
            page.click("#btn_search", timeout=20000)
            page.wait_for_function(
                """() => {
                    const info = $('#dataInput_info').text().trim();
                    return !$('#loadingDialog').is(':visible')
                        && !$('#modal-gagal').is(':visible')
                        && info
                        && !info.includes('Silakan pilih filter')
                        && !info.includes('NaN');
                }""",
                timeout=180000,
            )
            info_text = page.locator("#dataInput_info").inner_text().strip()
            total = parse_datatable_total(info_text)
            print(f"DigiPOS search {search_attempt}/3: {info_text}; total={total}", flush=True)
            if server_busy:
                raise RuntimeError(f"DigiPOS_SERVER_BUSY: {server_busy[-1]}")
            if total > 0:
                return total
            page.reload(wait_until="networkidle", timeout=45000)
        raise ValueError(f"DigiPOS showed zero entries for required date {report_date}")
    except PlaywrightTimeoutError as exc:
        if server_busy:
            raise RuntimeError(f"DigiPOS_SERVER_BUSY: {server_busy[-1]}") from exc
        raise RuntimeError(f"DigiPOS search timed out for {report_date}: {exc}") from exc
    finally:
        page.remove_listener("response", watch_response)


def _download_via_datatables_xhr(page, report_date: date, expected_total: int, output_path: Path) -> int:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

    if expected_total > MAX_TABLE_ROWS:
        raise ValueError(
            f"DigiPOS returned {expected_total} rows, beyond the known per-query "
            f"limit of {MAX_TABLE_ROWS}; split the report date window before processing"
        )

    route_state = {"rewritten": 0, "errors": [], "requests": []}
    endpoint_glob = "**/deposit/monitoring-riwayat-data*"

    def rewrite_length(route):
        request = route.request
        if ENDPOINT_PATH not in request.url or request.method != "POST":
            route.continue_()
            return
        try:
            post_data = _replace_datatables_length(
                request.post_data or "",
                expected_total,
            )
        except ValueError as exc:
            route_state["errors"].append(str(exc))
            route.continue_()
            return

        headers = {
            key: value
            for key, value in request.headers.items()
            if key.lower() != "content-length"
        }
        headers["content-length"] = str(len(post_data.encode("utf-8")))
        route_state["rewritten"] += 1
        route_state["requests"].append(request)
        route.continue_(post_data=post_data, headers=headers)

    page.route(endpoint_glob, rewrite_length)
    try:
        try:
            with page.expect_response(
                lambda response: (
                    ENDPOINT_PATH in response.url
                    and response.request.method == "POST"
                    and response.request in route_state["requests"]
                ),
                timeout=180000,
            ) as response_info:
                page.evaluate(
                    "length => $('#dataInput').DataTable().page.len(length).draw();",
                    expected_total,
                )
        except PlaywrightTimeoutError as exc:
            if route_state["errors"]:
                raise RuntimeError(route_state["errors"][-1]) from exc
            raise RuntimeError("DigiPOS XHR fallback timed out") from exc

        if route_state["errors"]:
            raise RuntimeError(route_state["errors"][-1])
        if not route_state["rewritten"]:
            raise RuntimeError("DigiPOS XHR fallback did not rewrite a DataTables request")

        response = response_info.value
        if response.status != 200:
            raise RuntimeError(f"DigiPOS XHR fallback returned HTTP {response.status}")
        rows = _rows_from_datatables_response(response.json(), expected_total)
        _write_xhr_export(output_path, rows)
        print(
            f"DigiPOS XHR fallback validated {len(rows)} rows for {report_date}",
            flush=True,
        )
        return len(rows)
    finally:
        page.unroute(endpoint_glob)


def _download_one_cluster(browser, cluster_id: str, report_date: date, password: str, output_dir: Path, max_attempts: int) -> dict:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

    username = f"{cluster_id}_A"
    canonical_filename = expected_source_filename(cluster_id, report_date.isoformat())
    artifact_name = f"{cluster_id}.xlsx"
    final_path = output_dir / artifact_name
    temp_path = output_dir / f".{cluster_id}.pending.xlsx"
    last_error = None

    for attempt in range(1, max_attempts + 1):
        page = None
        try:
            final_path.unlink(missing_ok=True)
            temp_path.unlink(missing_ok=True)
            print(f"DigiPOS cluster {cluster_id}, attempt {attempt}/{max_attempts}", flush=True)
            page = browser.new_page(accept_downloads=True, viewport={"width": 1280, "height": 800})
            page.set_default_timeout(20000)
            page.goto(f"{BASE_URL}/login", wait_until="networkidle", timeout=45000)
            page.evaluate(
                "([u, p]) => { $('#username').val(u).trigger('change'); $('#password').val(p).trigger('change'); }",
                [username, password],
            )
            page.evaluate("$('button:contains(\"Sign in\")').trigger('click')")
            page.wait_for_url("**/home", timeout=30000)
            table_total = _table_total_for_date(page, report_date)
            if page.locator(".buttons-excel").count():
                page.evaluate("$('#dataInput').DataTable().page.len(-1).draw()")
                page.wait_for_function(
                    "() => { const t = $('#dataInput').DataTable(); return !t.processing || !t.processing(); }",
                    timeout=60000,
                )
                page.wait_for_load_state("networkidle", timeout=45000)
                try:
                    with page.expect_download(timeout=90000) as download_info:
                        page.evaluate("$('.buttons-excel').trigger('click')")
                    download_info.value.save_as(str(temp_path))
                except PlaywrightTimeoutError:
                    print(
                        f"DigiPOS Excel download unavailable for {cluster_id}; "
                        "trying DataTables XHR fallback",
                        flush=True,
                    )
                    _download_via_datatables_xhr(
                        page,
                        report_date,
                        table_total,
                        temp_path,
                    )
            else:
                print(
                    f"DigiPOS Excel button missing for {cluster_id}; "
                    "using DataTables XHR fallback",
                    flush=True,
                )
                _download_via_datatables_xhr(
                    page,
                    report_date,
                    table_total,
                    temp_path,
                )

            row_count = validate_export_rows(
                _read_xlsx_rows(temp_path),
                report_date.isoformat(),
                table_total,
            )
            os.replace(temp_path, final_path)
            digest = hashlib.sha256(final_path.read_bytes()).hexdigest()
            print(f"DigiPOS validated {cluster_id}: {row_count} rows", flush=True)
            return {
                "cluster_id": cluster_id,
                "username": username,
                "filename": canonical_filename,
                "artifact": artifact_name,
                "row_count": row_count,
                "datatable_count": table_total,
                "sha256": digest,
            }
        except Exception as exc:
            last_error = exc
            print(f"DigiPOS export failed for {cluster_id}: {type(exc).__name__}", flush=True)
            temp_path.unlink(missing_ok=True)
            final_path.unlink(missing_ok=True)
            if attempt < max_attempts:
                time.sleep(3 * attempt)
        finally:
            if page:
                try:
                    page.close()
                except Exception:
                    pass
    raise RuntimeError(f"DigiPOS failed cluster {cluster_id} after {max_attempts} attempts: {last_error}")


def _download_cluster(cluster_id: str, report_date: date, password: str, output_dir: Path, max_attempts: int) -> dict:
    from playwright.sync_api import sync_playwright

    # Playwright's sync objects must be created, used, and closed on one thread.
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            return _download_one_cluster(
                browser, cluster_id, report_date, password, output_dir, max_attempts
            )
        finally:
            browser.close()


def scrape_all_clusters(
    report_date: str,
    password: str,
    output_dir: str = ".",
    max_attempts: int = 5,
    max_workers: int = 3,
) -> dict:
    """Scrape clusters concurrently; return only after all six exports validate."""
    parsed_date = parse_report_date(report_date)
    if not password:
        raise ValueError("DIGIPOS_PASSWORD Kestra secret is required")
    if max_attempts < 1:
        raise ValueError("max_attempts must be positive")
    if type(max_workers) is not int or not 1 <= max_workers <= len(CLUSTERS):
        raise ValueError(f"max_workers must be an integer between 1 and {len(CLUSTERS)}")
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    for cluster_id in CLUSTERS:
        (directory / f"{cluster_id}.xlsx").unlink(missing_ok=True)
        (directory / f".{cluster_id}.pending.xlsx").unlink(missing_ok=True)
    (directory / "digipos_manifest.json").unlink(missing_ok=True)

    print(f"DigiPOS acquiring {len(CLUSTERS)} clusters with {max_workers} workers", flush=True)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [
            executor.submit(
                _download_cluster, cluster_id, parsed_date, password, directory, max_attempts
            )
            for cluster_id in CLUSTERS
        ]
        files = [future.result() for future in futures]

    manifest = {
        "report_date": parsed_date.isoformat(),
        "source": "DigiPOS CMS Riwayat Saldo Transaksi",
        "all_clusters_validated": len(files) == len(CLUSTERS),
        "files": files,
    }
    (directory / "digipos_manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )
    return manifest
